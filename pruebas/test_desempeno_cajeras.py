"""Reportes (Karla, 2026-10-26): separador de miles, gráficos de venta por día y por hora, y el informe
Desempeño de cajeras con semáforo cada 100 comandas."""
from datetime import date, time
from decimal import Decimal

import openpyxl
import pytest
from django.urls import reverse

from nucleo.desempeno_cajeras_servicio import color_semaforo, desempeno, generar_excel
from nucleo.models import (
    ArqueoEfectivo, Cajera, CuentaContable, DiaArqueo, EventoAuditoria, Excepcion, PagoDeclarado, Parametro,
    Reclasificacion, RegularizacionDocumento, Venta,
)
from nucleo.templatetags.formato import miles
from nucleo.ventas_dashboard_servicio import grafico_por_dia, grafico_por_hora, resumen_periodo
from pruebas.perfiles_de_prueba import GRUPO_ENCARGADO, GRUPO_JEFE, GRUPO_SOLO_LECTURA, crear_usuario
from pruebas.test_ventas_dashboard import _crear_dia

D1, D2 = date(2026, 10, 3), date(2026, 10, 4)  # sábado y domingo


@pytest.fixture
def usuario(django_user_model):
    return crear_usuario(django_user_model, "operador", GRUPO_ENCARGADO, password="x")


# ------------------------------------------------------------------ formato y gráficos

def test_miles_redondea_y_separa():
    assert miles(1234567) == "1.234.567"
    assert miles(1234.6) == "1.235"
    assert miles(Decimal("-10688")) == "-10.688"
    assert miles(None) == ""


def test_grafico_por_dia_marca_fin_de_semana_pendientes_promedio_y_extremos():
    por_dia = {date(2026, 10, 1): 100.0, date(2026, 10, 2): 300.0}
    g = grafico_por_dia(date(2026, 10, 1), date(2026, 10, 31), por_dia, {date(2026, 10, 3): 50}, hoy=date(2026, 10, 4))
    assert [b["dia"] for b in g["barras"]] == [1, 2, 3, 4]  # hasta hoy, no hasta fin de mes
    assert g["barras"][2]["fin_de_semana"] and g["barras"][2]["pendiente"]
    assert g["barras"][1]["alto"] == 100 and g["promedio"] == 200 and g["alto_promedio"] == 67
    assert g["mejor"]["fecha"] == date(2026, 10, 2) and g["peor"]["fecha"] == date(2026, 10, 1)
    assert grafico_por_dia(date(2026, 1, 1), date(2026, 6, 30), por_dia, {}, hoy=date(2026, 7, 1)) is None


def test_grafico_por_hora_rellena_horas_y_marca_el_peak():
    g = grafico_por_hora({10: 100.0, 13: 300.0}, {10: 1, 13: 2})
    assert [b["hora"] for b in g["barras"]] == [10, 11, 12, 13]
    assert g["peak"]["hora"] == 13 and g["peak"]["pct"] == 75
    assert grafico_por_hora({}, {}) is None


def test_tablero_muestra_miles_unidades_por_comanda_y_dias_dentro_de_umbral(client, db, usuario):
    dia = _crear_dia(date(2026, 10, 1), DiaArqueo.Estado.CERRADO, usuario,
                     [("CAJA 1", 1250000, time(12, 0), [("EMPANADA", 3, 2500)]),
                      ("CAJA 1", 10000, time(13, 0), [("EMPANADA", 1, 2500)])])
    ArqueoEfectivo.objects.create(dia=dia, caja="CAJA 1", contado=1000, diferencia=-3000)
    r = resumen_periodo(date(2026, 10, 1), date(2026, 10, 31))
    assert r["unidades_por_comanda"] == 2
    assert r["control"]["dias_dentro_de_umbral"] == 1 and r["control"]["dias_cerrados_sin_diferencia"] == 0
    client.force_login(usuario)
    html = client.get(reverse("ventas_dashboard") + "?atajo=personalizado&desde=2026-10-01&hasta=2026-10-31").content.decode()
    assert "$1.260.000" in html and "2,0 unidades por comanda" in html
    assert "Días dentro de umbral" in html and "Venta por mes" not in html  # un solo mes: no se grafica
    assert "Hora peak" in html and "Promedio por día cerrado" in html


# ------------------------------------------------------------------ desempeño de cajeras

@pytest.fixture
def periodo(db, usuario):
    ana, bea = Cajera.objects.create(nombre="Ana"), Cajera.objects.create(nombre="Bea")
    d1 = _crear_dia(D1, DiaArqueo.Estado.CERRADO, usuario,
                    [("CAJA 1", 1000, time(10), [])] * 1 + [("CAJA 1", 2000, time(11), [])] + [("CAJA 2", 500, time(12), [])])
    d2 = _crear_dia(D2, DiaArqueo.Estado.EXPORTADO, usuario, [("CAJA 1", 700, time(10), []), ("CAJA 3", 900, time(10), [])])
    _crear_dia(date(2026, 10, 5), DiaArqueo.Estado.CONCILIADO_CON_EXCEPCIONES, usuario, [("CAJA 1", 99999, time(10), [])])
    ArqueoEfectivo.objects.create(dia=d1, caja="CAJA 1", cajera=ana, contado=0, diferencia=-8000)   # fuera de umbral
    ArqueoEfectivo.objects.create(dia=d1, caja="CAJA 2", cajera=bea, contado=0, diferencia=200)     # dentro
    ArqueoEfectivo.objects.create(dia=d2, caja="CAJA 1", cajera=bea, contado=0, diferencia=0)
    # CAJA 3 del día 2 no tiene arqueo: su comanda queda «sin cajera asignada».

    pago_ana = PagoDeclarado.objects.filter(venta__dia=d1, venta__caja="CAJA 1").first()
    for _ in range(2):  # dos reclasificaciones del mismo pago cuentan una vez
        Reclasificacion.objects.create(dia=d1, pago=pago_ana, medio_original="debito", medio_corregido="efectivo",
                                       monto=pago_ana.monto, origen="manual", regla="correccion_manual_de_medio", motivo="x")
    perdida, _ = CuentaContable.objects.get_or_create(codigo="6501004", defaults={"nombre": "Pérdida", "tipo": "perdida"})
    pago_bea = PagoDeclarado.objects.get(venta__dia=d2, venta__caja="CAJA 1")
    Excepcion.objects.create(dia=d2, tipo="diferencia_monto", pago=pago_bea, abierta=False, resolucion="perdida_ganancia",
                             cuenta_contable=perdida, monto_perdida_ganancia=300, comentario="cobró 400 de 700")
    RegularizacionDocumento.objects.create(dia=d1, tipo="CRED", codigo_original="CRE", n_interno=1, folio="F1",
                                           caja="CAJA 2", monto=500, forma_pago="efectivo")
    return {"ana": ana, "bea": bea}


def _fila(datos, nombre):
    return next(f for f in datos["filas"] if f["cajera"] == nombre)


def test_atribuye_errores_a_la_cajera_de_la_caja(periodo):
    datos = desempeno(date(2026, 10, 1), date(2026, 10, 31))
    ana, bea = _fila(datos, "Ana"), _fila(datos, "Bea")
    assert (ana["comandas"], ana["dias_trabajados"], ana["faltante"], ana["dias_fuera_umbral"], ana["errores_medio"]) == (2, 1, 8000, 1, 1)
    assert ana["errores"] == 2 and ana["tasa"] == 100 and ana["semaforo"] == "rojo"
    assert ana["reintegro_pendiente"] == 8000
    assert (bea["comandas"], bea["dias_trabajados"], bea["sobrante"], bea["cobros_de_menos"], bea["notas_credito"]) == (2, 2, 200, 1, 1)
    assert bea["errores"] == 1  # la NC se muestra pero no es error
    assert datos["sin_cajera"]["comandas"] == 1
    assert datos["dias_sin_cerrar"] == 1 and datos["totales"]["comandas"] == 5  # el día sin cerrar no cuenta
    assert [f["cajera"] for f in datos["top"]] == ["Ana", "Bea"]
    assert {i["tipo"] for i in datos["incidentes"]} == {"Faltante de efectivo", "Error de medio de pago", "Cobro de menos"}


def test_semaforo_configurable():
    assert color_semaforo(1, Decimal(1), Decimal(3)) == "verde"
    assert color_semaforo(2.5, Decimal(1), Decimal(3)) == "amarillo"
    assert color_semaforo(3.01, Decimal(1), Decimal(3)) == "rojo"
    assert color_semaforo(None, Decimal(1), Decimal(3)) == "gris"


def test_limites_del_semaforo_desde_configuracion(periodo):
    Parametro.objects.create(clave="semaforo_cajeras_verde", valor=100)
    Parametro.objects.create(clave="semaforo_cajeras_amarillo", valor=100)
    assert _fila(desempeno(date(2026, 10, 1), date(2026, 10, 31)), "Ana")["semaforo"] == "verde"


def test_pantalla_la_ven_todos_y_el_excel_solo_contabilidad(client, periodo, django_user_model, usuario):
    url = reverse("desempeno_cajeras") + "?desde=2026-10-01&hasta=2026-10-31"
    descarga = reverse("descargar_desempeno_cajeras") + "?desde=2026-10-01&hasta=2026-10-31"
    lectura = crear_usuario(django_user_model, "lector", GRUPO_SOLO_LECTURA)
    for u in (usuario, lectura):
        client.force_login(u)
        html = client.get(url).content.decode()
        assert "Detalle por cajera" in html and "Ana" in html and "$8.000" in html
        assert "Descargar Excel" not in html
        assert client.get(descarga).status_code in (302, 403)
    jefe = crear_usuario(django_user_model, "jefe", GRUPO_JEFE)
    client.force_login(jefe)
    assert "Descargar Excel" in client.get(url).content.decode()
    resp = client.get(descarga)
    assert resp.status_code == 200
    assert EventoAuditoria.objects.filter(accion="descargar_archivo", entidad_id="desempeno_cajeras").exists()


def test_excel_trae_resumen_e_incidentes(periodo, tmp_path):
    ruta = generar_excel(date(2026, 10, 1), date(2026, 10, 31), carpeta=tmp_path)
    wb = openpyxl.load_workbook(ruta)
    filas = list(wb["Resumen"].iter_rows(min_row=5, values_only=True))
    assert filas[0][0] == "Ana" and filas[0][12] == "Rojo" and filas[0][11] == 100
    assert wb["Incidentes"].max_row == 4  # encabezado + 3 incidentes


# ------------------------------------------------------------------ venta por cajera

def test_venta_por_cajera_ordenada_por_venta_por_dia_trabajado(periodo):
    from nucleo.desempeno_cajeras_servicio import venta_por_cajera
    v = venta_por_cajera(date(2026, 10, 1), date(2026, 10, 31))
    assert [f["cajera"] for f in v["filas"]] == ["Ana", "Bea", "Sin cajera asignada"]
    ana, bea = v["filas"][0], v["filas"][1]
    assert (ana["venta"], ana["comandas"], ana["ticket"], ana["venta_por_dia"]) == (3000, 2, 1500, 3000)
    assert (bea["venta"], bea["dias_trabajados"], bea["ticket"], bea["venta_por_dia"]) == (1200, 2, 600, 600)
    assert v["ticket_general"] == 1020 and v["lider"]["cajera"] == "Ana"
    assert round(ana["ticket_vs_general"], 1) == 47.1


def test_venta_por_cajera_en_tablero_informe_y_excel(client, periodo, usuario, tmp_path):
    client.force_login(usuario)
    for url in (reverse("ventas_dashboard") + "?atajo=personalizado&desde=2026-10-01&hasta=2026-10-31",
                reverse("desempeno_cajeras") + "?desde=2026-10-01&hasta=2026-10-31"):
        html = client.get(url).content.decode()
        assert "Venta por cajera" in html and "Lidera" in html and "$1.500" in html
    wb = openpyxl.load_workbook(generar_excel(date(2026, 10, 1), date(2026, 10, 31), carpeta=tmp_path))
    hoja = list(wb["Venta por cajera"].iter_rows(min_row=2, values_only=True))
    assert hoja[0][0] == "Ana" and hoja[0][5] == 1500 and hoja[0][6] == 3000
