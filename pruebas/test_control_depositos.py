"""Control de Depósitos (Karla, 2026-09-21/22): un segundo registro de
depósito, más completo que el simple de Arqueo de efectivo, que nace al
cerrar el día y debe pasar por la Analista de Ingresos (lo completa) y el
Jefe de Contabilidad (lo aprueba u observa, con evidencia real)."""
from datetime import date, time
from decimal import Decimal

import pytest
from pruebas.perfiles_de_prueba import GRUPO_ENCARGADO, GRUPO_GERENTE, crear_usuario
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from nucleo.arqueo_servicio import registrar_conteo
from nucleo.permisos import PermisoDenegado
from nucleo.cierre_servicio import aprobar_y_cerrar_dia
from nucleo.configuracion_servicio import actualizar_umbral
from nucleo.depositos_servicio import (
    ControlDepositoInvalido, agregar_deposito_adicional, crear_al_cerrar, depositos_con_diferencia_del_mes,
    generar_captura_deposito, generar_lineas_captura_deposito, generar_reporte_mensual, registrar_detalle,
    resumen_por_estado, validar_deposito,
)
from nucleo.models import (
    ArqueoEfectivo, Deposito, DepositoBancario, DiaArqueo, EventoAuditoria, LineaVenta, LoteCarga, PagoDeclarado, Venta,
)

FECHA = date(2026, 9, 1)


def _comprobante(nombre="comprobante.pdf"):
    return SimpleUploadedFile(nombre, b"%PDF-1.4 contenido", content_type="application/pdf")


@pytest.fixture
def usuario(django_user_model):
    return crear_usuario(django_user_model, "operador", GRUPO_GERENTE)


@pytest.fixture
def gerente(django_user_model):
    from django.contrib.auth.models import Group
    from nucleo.documentos_servicio import GRUPO_GERENTE_ADM_FINANZAS
    u = django_user_model.objects.create_user(username="gerente", password="x")
    u.groups.add(Group.objects.get_or_create(name=GRUPO_GERENTE_ADM_FINANZAS)[0])
    return u


@pytest.fixture
def encargada(django_user_model):
    return crear_usuario(django_user_model, "encargada", GRUPO_ENCARGADO)


@pytest.fixture
def jefe(django_user_model):
    from django.contrib.auth.models import Group
    from nucleo.documentos_servicio import GRUPO_JEFE_CONTABILIDAD
    u = django_user_model.objects.create_user(username="jefe", password="x")
    u.groups.add(Group.objects.get_or_create(name=GRUPO_JEFE_CONTABILIDAD)[0])
    return u


@pytest.fixture
def dia_cerrado(db, usuario):
    dia = DiaArqueo.objects.create(fecha=FECHA, operado_por=usuario)
    lote = LoteCarga.objects.create(
        dia=dia, fuente=LoteCarga.Fuente.POS, nombre_archivo="x", huella_sha256="a",
        ruta_almacenamiento="/x", fecha_declarada=FECHA, cargado_por=usuario,
    )
    venta = Venta.objects.create(
        dia=dia, lote=lote, n_interno=1, nro_comprobante="B0000001", tipo_doc="FAC", nro_comanda="1",
        estado="COBRADA", caja="CAJA 1", turno="T", hora_cierre=time(10, 0), total_comanda=50000,
        cliente="CLIENTE OCASIONAL",
    )
    LineaVenta.objects.create(venta=venta, articulo="X", cantidad=1, precio_unitario=50000, total_linea=50000)
    PagoDeclarado.objects.create(venta=venta, medio="efectivo", monto=50000)
    arqueo = ArqueoEfectivo.objects.create(dia=dia, caja="CAJA 1")
    registrar_conteo(arqueo, fondo_inicial=Decimal(0), gastos_autorizados=Decimal(0),
                      total_contado=Decimal(50000), usuario=usuario)
    aprobar_y_cerrar_dia(dia, usuario, motivo_excepcion="prueba")
    return dia


def test_el_deposito_bancario_nace_solo_al_cerrar_el_dia(dia_cerrado):
    depositos = DepositoBancario.objects.filter(dia=dia_cerrado)
    assert depositos.count() == 1
    deposito = depositos.first()
    assert deposito.estado == DepositoBancario.Estado.PENDIENTE_DEPOSITO
    assert deposito.monto_esperado == 50000
    assert deposito.monto_depositado is None and deposito.diferencia is None


def test_reabrir_y_cerrar_de_nuevo_no_duplica_el_deposito_pendiente(dia_cerrado, usuario):
    from nucleo.cierre_servicio import reabrir_dia
    reabrir_dia(dia_cerrado, usuario, motivo="corrección")
    aprobar_y_cerrar_dia(dia_cerrado, usuario, motivo_excepcion="prueba")
    assert DepositoBancario.objects.filter(dia=dia_cerrado).count() == 1


def test_el_deposito_convive_con_el_simple_de_arqueo_de_efectivo(dia_cerrado, usuario):
    Deposito.objects.create(dia=dia_cerrado, monto=50000, comprobante="123", fecha=FECHA)
    assert Deposito.objects.filter(dia=dia_cerrado).exists()
    assert DepositoBancario.objects.filter(dia=dia_cerrado).exists()  # no se reemplazan entre sí


# ---------------------------------------------------------- registrar_detalle

def test_la_analista_registra_el_detalle_y_queda_pendiente_de_validacion(dia_cerrado, usuario):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)

    registrar_detalle(
        deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="12345678",
        monto_depositado=Decimal(49999), numero_operacion="OP-001", responsable_traslado="Rosa Díaz",
        comprobante=_comprobante(), usuario=usuario,
    )
    deposito.refresh_from_db()
    assert deposito.estado == DepositoBancario.Estado.PENDIENTE_VALIDACION
    assert deposito.diferencia == -1  # faltó $1 por depositar
    assert deposito.comprobante_archivo and deposito.comprobante_nombre_original == "comprobante.pdf"
    assert deposito.responsable_traslado == "Rosa Díaz"
    from nucleo.models import EventoAuditoria
    assert EventoAuditoria.objects.filter(accion="registrar_detalle_deposito").exists()


@pytest.mark.parametrize("campos,mensaje", [
    (dict(banco="", cuenta_bancaria="1", monto_depositado=Decimal(1), numero_operacion="OP1", responsable_traslado="Rosa"), "banco"),
    (dict(banco="BCI", cuenta_bancaria="1", monto_depositado=None, numero_operacion="OP1", responsable_traslado="Rosa"), "monto"),
    (dict(banco="BCI", cuenta_bancaria="1", monto_depositado=Decimal(1), numero_operacion="", responsable_traslado="Rosa"), "operación"),
    (dict(banco="BCI", cuenta_bancaria="1", monto_depositado=Decimal(1), numero_operacion="OP1", responsable_traslado=""), "retiró"),
])
def test_faltan_datos_obligatorios_para_registrar_el_deposito(dia_cerrado, usuario, campos, mensaje):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    with pytest.raises(ControlDepositoInvalido, match=mensaje):
        registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), comprobante=_comprobante(), usuario=usuario, **campos)


def test_el_comprobante_es_obligatorio(dia_cerrado, usuario):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    with pytest.raises(ControlDepositoInvalido, match="[Cc]omprobante"):
        registrar_detalle(
            deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
            monto_depositado=Decimal(50000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
            comprobante=None, usuario=usuario,
        )


def test_un_deposito_ya_revisado_no_se_puede_editar(dia_cerrado, usuario, jefe):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(50000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    validar_deposito(deposito, aprobar=True, movimiento_bancario_referencia="Cartola BCI folio 1",
                      comentarios="", usuario=jefe)
    with pytest.raises(ControlDepositoInvalido, match="ya fue revisado"):
        registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                           monto_depositado=Decimal(50000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                           comprobante=_comprobante(), usuario=usuario)


# ------------------------------------------------------------- validar_deposito

@pytest.fixture
def deposito_pendiente_de_validar(dia_cerrado, usuario):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(49999), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    return deposito


def test_aprobar_exige_la_referencia_del_movimiento_bancario(deposito_pendiente_de_validar, jefe):
    with pytest.raises(ControlDepositoInvalido, match="evidencia"):
        validar_deposito(deposito_pendiente_de_validar, aprobar=True, movimiento_bancario_referencia="",
                          comentarios="", usuario=jefe)


def test_observar_no_exige_la_referencia_pero_la_diferencia_queda_visible(deposito_pendiente_de_validar, jefe):
    validar_deposito(deposito_pendiente_de_validar, aprobar=False, movimiento_bancario_referencia="",
                      comentarios="Falta cuadrar con la cajera", usuario=jefe)
    deposito_pendiente_de_validar.refresh_from_db()
    assert deposito_pendiente_de_validar.estado == DepositoBancario.Estado.OBSERVADO
    assert deposito_pendiente_de_validar.diferencia == -1  # nunca se compensa ni se borra sola
    assert deposito_pendiente_de_validar.resultado_cotejo == "con_diferencia"


def test_aprobar_con_evidencia_deja_todo_trazado(deposito_pendiente_de_validar, jefe):
    validar_deposito(
        deposito_pendiente_de_validar, aprobar=True,
        movimiento_bancario_referencia="Abono visto en cartola BCI del 02/09, folio 445566",
        comentarios="Diferencia de $1 por redondeo, aceptable", usuario=jefe,
    )
    deposito_pendiente_de_validar.refresh_from_db()
    assert deposito_pendiente_de_validar.estado == DepositoBancario.Estado.APROBADO
    assert deposito_pendiente_de_validar.aprobado_por == jefe
    assert deposito_pendiente_de_validar.movimiento_bancario_referencia.startswith("Abono visto")
    from nucleo.models import EventoAuditoria
    evento = EventoAuditoria.objects.get(accion="validar_deposito")
    assert evento.usuario == jefe and evento.valor_posterior["estado"] == "aprobado"


def test_una_encargada_no_puede_validar(deposito_pendiente_de_validar, encargada):
    with pytest.raises(PermisoDenegado):
        validar_deposito(deposito_pendiente_de_validar, aprobar=True, movimiento_bancario_referencia="x",
                          comentarios="", usuario=encargada)


def test_no_se_puede_validar_dos_veces(deposito_pendiente_de_validar, jefe):
    validar_deposito(deposito_pendiente_de_validar, aprobar=True, movimiento_bancario_referencia="x",
                      comentarios="", usuario=jefe)
    with pytest.raises(ControlDepositoInvalido, match="pendiente de validación"):
        validar_deposito(deposito_pendiente_de_validar, aprobar=True, movimiento_bancario_referencia="x",
                          comentarios="", usuario=jefe)


def test_una_diferencia_sobre_el_umbral_no_se_puede_aprobar_solo_observar(dia_cerrado, usuario, jefe):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(20000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)  # esperado 50.000, faltan 30.000
    with pytest.raises(ControlDepositoInvalido, match="umbral"):
        validar_deposito(deposito, aprobar=True, movimiento_bancario_referencia="x", comentarios="", usuario=jefe)
    validar_deposito(deposito, aprobar=False, movimiento_bancario_referencia="", comentarios="revisar", usuario=jefe)
    deposito.refresh_from_db()
    assert deposito.estado == DepositoBancario.Estado.OBSERVADO


def test_el_umbral_del_deposito_es_configurable(dia_cerrado, usuario, jefe):
    actualizar_umbral("umbral_diferencia_deposito", "100000", usuario)
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(20000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    validar_deposito(deposito, aprobar=True, movimiento_bancario_referencia="x", comentarios="", usuario=jefe)
    deposito.refresh_from_db()
    assert deposito.estado == DepositoBancario.Estado.APROBADO


def test_resumen_por_estado_cuenta_bien(dia_cerrado, usuario, jefe):
    otro_dia = DiaArqueo.objects.create(fecha=date(2026, 9, 2), operado_por=usuario)
    DepositoBancario.objects.create(dia=otro_dia, monto_esperado=1000, registrado_por=usuario)
    todos = DepositoBancario.objects.all()
    resumen = resumen_por_estado(todos)
    assert resumen["pendiente_deposito"] == 2
    assert resumen["sin_cotejar"] == 2


# ------------------------------------------------------------ más de un depósito

def test_se_puede_agregar_mas_de_un_deposito_para_el_mismo_dia(dia_cerrado, usuario):
    agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(20000), usuario=usuario)
    assert DepositoBancario.objects.filter(dia=dia_cerrado).count() == 2


def test_el_adicional_reparte_el_total_del_dia_no_se_suma_encima(dia_cerrado, usuario):
    """Día 05 real: el adicional de $50.000 se sumaba sobre lo contado y dejaba una diferencia falsa."""
    principal = DepositoBancario.objects.get(dia=dia_cerrado)
    assert principal.monto_esperado == 50000
    adicional = agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(20000), usuario=usuario)
    principal.refresh_from_db()
    assert (principal.monto_esperado, adicional.monto_esperado) == (30000, 20000)
    assert sum(d.monto_esperado for d in DepositoBancario.objects.filter(dia=dia_cerrado)) == 50000  # = lo contado
    assert EventoAuditoria.objects.filter(accion="agregar_deposito_adicional", entidad_id=str(adicional.pk)).exists()


def test_el_adicional_no_puede_ser_igual_o_mayor_a_lo_que_queda_por_depositar(dia_cerrado, usuario):
    with pytest.raises(ControlDepositoInvalido, match="se reparte el total"):
        agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(50000), usuario=usuario)
    with pytest.raises(ControlDepositoInvalido, match="se reparte el total"):
        agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(60000), usuario=usuario)
    assert DepositoBancario.objects.filter(dia=dia_cerrado).count() == 1


def test_no_se_toca_un_deposito_ya_revisado_y_si_todos_lo_estan_se_avisa(dia_cerrado, usuario, jefe):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(50000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    validar_deposito(deposito, aprobar=True, movimiento_bancario_referencia="Cartola", comentarios="", usuario=jefe)
    with pytest.raises(ControlDepositoInvalido, match="ya fueron revisados"):
        agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(1000), usuario=usuario)
    deposito.refresh_from_db()
    assert deposito.monto_esperado == 50000 and deposito.diferencia == 0


def test_repartido_el_total_cuadra_sin_diferencia_falsa(dia_cerrado, usuario, banco_estado):
    agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(20000), usuario=usuario)
    principal, adicional = DepositoBancario.objects.filter(dia=dia_cerrado).order_by("pk")
    for d, monto, op in ((principal, 30000, "A"), (adicional, 20000, "B")):
        registrar_detalle(d, fecha_deposito=date(2026, 9, 2), banco="Banco Estado 10397-7", cuenta_bancaria="1",
                           monto_depositado=Decimal(monto), numero_operacion=op, responsable_traslado="Rosa Díaz",
                           comprobante=_comprobante(f"{op}.pdf"), usuario=usuario)
    assert [d.diferencia for d in DepositoBancario.objects.filter(dia=dia_cerrado)] == [0, 0]


def test_el_deposito_adicional_exige_un_monto_valido(dia_cerrado, usuario):
    with pytest.raises(ControlDepositoInvalido, match="monto"):
        agregar_deposito_adicional(dia_cerrado, monto_esperado=None, usuario=usuario)
    with pytest.raises(ControlDepositoInvalido, match="monto"):
        agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(0), usuario=usuario)


# --------------------------------------------------------- reporte mensual

def test_el_reporte_mensual_solo_lleva_aprobados_u_observados_con_diferencia(dia_cerrado, usuario, jefe):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(49999), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    validar_deposito(deposito, aprobar=True, movimiento_bancario_referencia="Cartola folio 1",
                      comentarios="ok", usuario=jefe)

    sin_diferencia = crear_al_cerrar(DiaArqueo.objects.create(fecha=date(2026, 9, 2), operado_por=usuario), usuario)
    sin_diferencia.monto_esperado = 0
    sin_diferencia.save()

    filas = depositos_con_diferencia_del_mes(2026, 9)
    assert filas == [deposito]


def test_generar_reporte_mensual_produce_un_excel(dia_cerrado, usuario, jefe, tmp_path):
    from openpyxl import load_workbook

    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(49999), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    validar_deposito(deposito, aprobar=True, movimiento_bancario_referencia="Cartola folio 1",
                      comentarios="ok", usuario=jefe)

    ruta = generar_reporte_mensual(2026, 9, tmp_path)
    filas = list(load_workbook(ruta).active.iter_rows(values_only=True))
    assert filas[0][0] == "Fecha del día"
    assert len(filas) == 4  # encabezado, el depósito, una línea en blanco y el total
    assert filas[1][8] == "Rosa Díaz"  # responsable del traslado
    assert filas[1][6] == "Sí"  # con diferencia
    assert filas[3][0] == "TOTAL"


# ------------------------------------------------------------------- pantalla

def test_la_pantalla_de_control_de_depositos_se_ve_y_permite_registrar(client, dia_cerrado, usuario):
    client.force_login(usuario)
    resp = client.get(reverse("control_depositos"))
    assert resp.status_code == 200
    html = resp.content.decode()
    assert "Pendiente de depósito" in html

    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    resp = client.post(reverse("registrar_detalle_deposito", args=[deposito.pk]), {
        "fecha_deposito": "2026-09-02", "banco": "BCI", "cuenta_bancaria": "12345",
        "monto_depositado": "50.000", "numero_operacion": "OP-9", "responsable_traslado": "Rosa Díaz",
        "comprobante": _comprobante(),
    })
    assert resp.status_code == 302
    deposito.refresh_from_db()
    assert deposito.estado == "pendiente_validacion" and deposito.monto_depositado == 50000


def test_solo_el_jefe_puede_aprobar_desde_la_pantalla(client, dia_cerrado, usuario, jefe, encargada):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(50000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)

    client.force_login(encargada)
    resp = client.post(reverse("validar_deposito", args=[deposito.pk]), {
        "accion": "aprobar", "movimiento_bancario_referencia": "x",
    })
    assert resp.status_code == 403
    deposito.refresh_from_db()
    assert deposito.estado == "pendiente_validacion"  # no pasó nada: la Encargada no valida

    client.force_login(jefe)
    client.post(reverse("validar_deposito", args=[deposito.pk]), {
        "accion": "aprobar", "movimiento_bancario_referencia": "Cartola BCI folio 2",
    })
    deposito.refresh_from_db()
    assert deposito.estado == "aprobado"


def test_agregar_deposito_adicional_desde_la_pantalla(client, dia_cerrado, usuario):
    client.force_login(usuario)
    resp = client.post(reverse("agregar_deposito_adicional", args=[dia_cerrado.pk]), {"monto_esperado": "20.000"})
    assert resp.status_code == 302
    assert DepositoBancario.objects.filter(dia=dia_cerrado).count() == 2


def test_descargar_reporte_mensual_desde_la_pantalla(client, dia_cerrado, usuario, jefe):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(49999), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    validar_deposito(deposito, aprobar=True, movimiento_bancario_referencia="Cartola folio 1",
                      comentarios="ok", usuario=jefe)
    client.force_login(usuario)
    resp = client.get(reverse("descargar_reporte_mensual_depositos"), {"anio": 2026, "mes": 9})
    assert resp.status_code == 200
    assert resp["Content-Disposition"].endswith('reporte_depositos_2026-09.xlsx"')


def test_el_menu_muestra_control_de_depositos(client, usuario, db):
    client.force_login(usuario)
    html = client.get(reverse("ventas_dashboard")).content.decode()
    assert reverse("control_depositos") in html


# ------------------------------------------------- reporte mensual (2026-09-24)

def test_el_reporte_mensual_incluye_los_depositos_sin_validar_con_diferencia(dia_cerrado, usuario, tmp_path):
    """El reporte salía vacío porque solo listaba depósitos ya validados con diferencia."""
    from openpyxl import load_workbook
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)
    registrar_detalle(deposito, fecha_deposito=date(2026, 9, 2), banco="BCI", cuenta_bancaria="1",
                       monto_depositado=Decimal(40000), numero_operacion="OP1", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    assert depositos_con_diferencia_del_mes(2026, 9) == [deposito]
    filas = list(load_workbook(generar_reporte_mensual(2026, 9, tmp_path)).active.iter_rows(values_only=True))
    assert filas[1][1] == "Pendiente de validación" and filas[1][6] == "Sí"


def test_el_reporte_mensual_de_un_mes_sin_depositos_lo_dice(db, tmp_path):
    from openpyxl import load_workbook
    filas = list(load_workbook(generar_reporte_mensual(2026, 1, tmp_path)).active.iter_rows(values_only=True))
    assert filas[1][0] == "Sin depósitos registrados en 01/2026"


def test_un_anio_o_mes_mal_escrito_no_rompe_la_pantalla_ni_el_reporte(client, usuario, db):
    client.force_login(usuario)
    assert client.get(reverse("control_depositos"), {"anio_reporte": "abc", "mes_reporte": "13"}).status_code == 200
    assert client.get(reverse("descargar_reporte_mensual_depositos"), {"anio": "x", "mes": "99"}).status_code == 200


# ------------------------------------------------ captura de depósito (Odoo)

@pytest.fixture
def banco_estado(db):
    from nucleo.models import CuentaContable
    return CuentaContable.objects.create(
        codigo="1102001", nombre="Banco Estado 10397-7", tipo="banco",
        diario_odoo="Banco Estado", metodo_pago_odoo="Pagos Entrantes (Banco Estado)",
    )


def _depositar(dia, usuario, banco="Banco Estado 10397-7", monto=50000, fecha_deposito=date(2026, 9, 24), op="OP1"):
    deposito = DepositoBancario.objects.filter(dia=dia).first()
    registrar_detalle(deposito, fecha_deposito=fecha_deposito, banco=banco, cuenta_bancaria="10397-7",
                       monto_depositado=Decimal(monto), numero_operacion=op, responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante(), usuario=usuario)
    return deposito


def test_la_captura_de_deposito_lleva_el_diario_y_metodo_del_banco_elegido(dia_cerrado, usuario, banco_estado, tmp_path):
    """Formato del ejemplo de Karla: Recibir / empresa / importe / fecha del depósito / cuenta efectivo /
    memo con el día de caja / diario y método de pago del banco."""
    from openpyxl import load_workbook
    deposito = _depositar(dia_cerrado, usuario)
    ruta = generar_captura_deposito(deposito, usuario, tmp_path)
    filas = list(load_workbook(ruta).active.iter_rows(values_only=True))
    assert filas[0] == (
        "Tipo de pago", "Cliente/proveedor", "Importe en moneda de la empresa firmado", "Fecha",
        "Cuenta Contable", "Memo", "Diario",
    )
    assert filas[1][:3] == ("Recibir", "AGRICOLA HACIENDA HUENTELAUQUEN LTDA", 50000)
    assert filas[1][3].date() == date(2026, 9, 24)
    assert filas[1][4:] == (
        "1101004 Efectivo Pos Resto", "Deposito Efectivo Caja del dia 01-09-2026", "Banco Estado",
    )
    deposito.refresh_from_db()
    assert deposito.captura_exportada_en is not None
    assert EventoAuditoria.objects.filter(accion="exportar_captura_deposito", entidad_id=str(deposito.pk)).exists()


def test_cada_banco_usa_su_propio_diario(dia_cerrado, usuario, banco_estado):
    from nucleo.models import CuentaContable
    CuentaContable.objects.create(
        codigo="1102002", nombre="BCI Cta Cte", tipo="banco", diario_odoo="BCI", metodo_pago_odoo="Pagos Entrantes (BCI)",
    )
    deposito = _depositar(dia_cerrado, usuario, banco="BCI Cta Cte")
    fila = generar_lineas_captura_deposito(deposito)[0]
    assert fila["Diario"] == "BCI" and "Método de pago" not in fila


def test_la_captura_esta_disponible_antes_de_la_validacion_del_jefe(dia_cerrado, usuario, banco_estado, tmp_path):
    deposito = _depositar(dia_cerrado, usuario)
    assert deposito.estado == "pendiente_validacion"
    assert generar_captura_deposito(deposito, usuario, tmp_path).exists()


def test_sin_datos_del_deposito_no_hay_captura(dia_cerrado, usuario, banco_estado):
    deposito = DepositoBancario.objects.get(dia=dia_cerrado)  # pendiente de depósito
    with pytest.raises(ControlDepositoInvalido, match="Primero ingresa"):
        generar_lineas_captura_deposito(deposito)


def test_un_banco_sin_diario_bloquea_la_captura_y_dice_donde_corregirlo(dia_cerrado, usuario, banco_estado):
    deposito = _depositar(dia_cerrado, usuario)
    banco_estado.diario_odoo = ""
    banco_estado.save()
    with pytest.raises(ControlDepositoInvalido, match="Diario.*Configuración"):
        generar_lineas_captura_deposito(deposito)
    banco_estado.diario_odoo, banco_estado.metodo_pago_odoo = "Banco Estado", ""
    banco_estado.save()
    assert generar_lineas_captura_deposito(deposito)[0]["Diario"] == "Banco Estado"  # el método de pago ya no se usa


def test_un_banco_que_no_esta_en_configuracion_bloquea_la_captura(dia_cerrado, usuario):
    deposito = _depositar(dia_cerrado, usuario, banco="Banco Fantasma")
    with pytest.raises(ControlDepositoInvalido, match="no existe entre las cuentas de Banco"):
        generar_lineas_captura_deposito(deposito)


def test_con_dos_depositos_el_mismo_dia_el_memo_lleva_la_operacion(dia_cerrado, usuario, banco_estado):
    primero = _depositar(dia_cerrado, usuario, monto=30000, op="OP-A")  # se espera 50.000; luego se reparte
    segundo = agregar_deposito_adicional(dia_cerrado, monto_esperado=Decimal(20000), usuario=usuario)
    registrar_detalle(segundo, fecha_deposito=date(2026, 9, 24), banco="Banco Estado 10397-7", cuenta_bancaria="1",
                       monto_depositado=Decimal(20000), numero_operacion="OP-B", responsable_traslado="Rosa Díaz",
                       comprobante=_comprobante("otro.pdf"), usuario=usuario)
    assert generar_lineas_captura_deposito(primero)[0]["Memo"].endswith("- Op. OP-A")
    assert generar_lineas_captura_deposito(segundo)[0]["Memo"].endswith("- Op. OP-B")


def test_boton_y_descarga_de_la_captura_desde_la_pantalla(client, dia_cerrado, usuario, banco_estado):
    deposito = _depositar(dia_cerrado, usuario)
    client.force_login(usuario)
    html = client.get(reverse("control_depositos"), {"deposito": deposito.pk}).content.decode()
    # Karla, 2026-09-24: el botón de descarga vive en «Exportaciones a Odoo»; aquí solo queda el enlace
    assert "Descargar captura de depósito" not in html and reverse("captura_deposito", args=[deposito.pk]) not in html
    assert reverse("exportaciones_odoo") in html and "Captura Depósito" in html
    exportaciones = client.get(reverse("exportaciones_odoo")).content.decode()
    assert reverse("generar_captura_deposito", args=[deposito.pk]) in exportaciones  # primero se genera
    # Karla, 2026-09-26: bajar un archivo (GET) ya no lo genera: primero se genera (POST) y después se descarga
    sin_generar = client.get(reverse("captura_deposito", args=[deposito.pk]), follow=True)
    assert sin_generar.redirect_chain and "todavía no se ha generado" in sin_generar.content.decode()
    generada = client.post(reverse("generar_captura_deposito", args=[deposito.pk]))
    assert generada.status_code == 302
    exportaciones = client.get(reverse("exportaciones_odoo")).content.decode()
    assert "Descargar Captura Depósito" in exportaciones and reverse("captura_deposito", args=[deposito.pk]) in exportaciones
    resp = client.get(reverse("captura_deposito", args=[deposito.pk]))
    assert resp.status_code == 200 and resp["Content-Disposition"].startswith("attachment")


def test_la_pantalla_avisa_si_al_banco_le_falta_el_diario(client, dia_cerrado, usuario, banco_estado):
    deposito = _depositar(dia_cerrado, usuario)
    banco_estado.diario_odoo = ""
    banco_estado.save()
    client.force_login(usuario)
    html = client.get(reverse("control_depositos"), {"deposito": deposito.pk}).content.decode()
    assert "pendiente:" in html and "le falta el Diario" in html and reverse("captura_deposito", args=[deposito.pk]) not in html
    exportaciones = client.get(reverse("exportaciones_odoo")).content.decode()
    assert "le falta el Diario" in exportaciones and reverse("captura_deposito", args=[deposito.pk]) not in exportaciones
    resp = client.get(reverse("captura_deposito", args=[deposito.pk]), follow=True)
    assert "le falta el Diario" in resp.content.decode()


def test_arqueo_de_efectivo_ya_no_tiene_el_formulario_de_deposito(client, dia_cerrado, usuario):
    client.force_login(usuario)
    html = client.get(f"/arqueo/{FECHA.isoformat()}/").content.decode()
    assert "Registrar depósito" not in html and "Control de Depósitos" in html
    assert client.post(f"/arqueo/{FECHA.isoformat()}/deposito/", {"monto": "1"}).status_code == 404
