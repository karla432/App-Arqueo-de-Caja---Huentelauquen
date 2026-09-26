"""Fase A (Karla, 2026-09-21): corrección y trazabilidad de la forma de pago.

Una venta declarada como Transbank que no aparece en el proveedor, con un
sobrante de efectivo por el mismo monto, se corrige sola a efectivo — y al
revés: una venta declarada como efectivo, con un cobro del proveedor sin
venta y un faltante por el mismo monto, se corrige a tarjeta. Todo queda en
`Reclasificacion` y en el reporte de auditoría, sin aprobación previa.
"""
from datetime import date, datetime, time
from decimal import Decimal

import pytest
from pruebas.perfiles_de_prueba import GRUPO_GERENTE, crear_usuario
from django.urls import reverse
from openpyxl import load_workbook

from nucleo.arqueo_servicio import calcular_ajustes_por_caja, registrar_conteo
from nucleo.auditoria_servicio import filas_reclasificacion, generar_archivo_auditoria
from nucleo.conciliacion_servicio import ejecutar_conciliacion
from nucleo.models import (
    ArqueoEfectivo, Coincidencia, DiaArqueo, EventoAuditoria, Excepcion, LineaVenta, LoteCarga, MotivoResolucion,
    MovimientoProveedor, PagoDeclarado, Reclasificacion, Venta,
)
from nucleo.resolucion_servicio import resolver_excepcion

FECHA = date(2026, 9, 1)


@pytest.fixture
def usuario(django_user_model):
    return crear_usuario(django_user_model, "operador", GRUPO_GERENTE)


@pytest.fixture
def dia(db, usuario):
    dia = DiaArqueo.objects.create(fecha=FECHA, operado_por=usuario)
    for fuente, nombre in ((LoteCarga.Fuente.POS, "pos"), (LoteCarga.Fuente.TRANSBANK, "tbk")):
        LoteCarga.objects.create(
            dia=dia, fuente=fuente, nombre_archivo=nombre, huella_sha256=nombre, ruta_almacenamiento="/x",
            fecha_declarada=FECHA, cargado_por=usuario,
        )
    return dia


def _venta(dia, n, medio, monto, hora, caja="CAJA 1"):
    venta = Venta.objects.create(
        dia=dia, lote=dia.lotes.get(fuente="pos"), n_interno=n, nro_comprobante=f"B{n:07d}", tipo_doc="FAC",
        nro_comanda=str(n), estado="COBRADA", caja=caja, turno="T", hora_cierre=hora, total_comanda=monto,
        cliente="CLIENTE OCASIONAL",
    )
    LineaVenta.objects.create(venta=venta, articulo="X", cantidad=1, precio_unitario=monto, total_linea=monto)
    return PagoDeclarado.objects.create(venta=venta, medio=medio, monto=monto)


def _cobro(dia, monto, hora, minuto=0, medio="debito", autorizacion="A1"):
    return MovimientoProveedor.objects.create(
        dia=dia, lote=dia.lotes.get(fuente="transbank"), proveedor="transbank", tipo="Venta", estado="PENDIENTE",
        fecha_hora=datetime(2026, 9, 1, hora, minuto), medio=medio, tipo_tarjeta=medio, monto_original=monto,
        monto_venta_valido_abono=monto, autorizacion=autorizacion, terminal="T1",
    )


def _contar(dia, caja, total):
    arqueo, _ = ArqueoEfectivo.objects.get_or_create(dia=dia, caja=caja)
    arqueo.contado = Decimal(total)
    arqueo.save()


# ------------------------------------------------- tarjeta -> efectivo

@pytest.fixture
def caja_con_sobrante(dia):
    """Efectivo $20.000 + una venta con tarjeta de $30.000 que no existe en
    Transbank + otra con tarjeta que sí. La caja tiene $30.000 de más."""
    _venta(dia, 1, "efectivo", 20000, time(10, 0))
    falsa = _venta(dia, 2, "debito", 30000, time(11, 0))
    _venta(dia, 3, "debito", 10000, time(12, 0))
    _cobro(dia, 10000, 12, 1, autorizacion="A3")
    _contar(dia, "CAJA 1", 50000)
    return dia, falsa


def test_venta_declarada_como_tarjeta_que_no_esta_en_el_proveedor_pasa_a_efectivo_si_sobra_ese_monto(caja_con_sobrante, usuario):
    dia, falsa = caja_con_sobrante
    ejecutar_conciliacion(dia, usuario)

    falsa.refresh_from_db()
    assert (falsa.medio, falsa.medio_aprobado) == ("debito", "efectivo")  # lo declarado no se toca
    r = Reclasificacion.objects.get(pago=falsa)
    assert (r.medio_original, r.medio_corregido, r.monto) == ("debito", "efectivo", 30000)
    assert r.origen == "automatica" and r.usuario is None
    assert r.regla == "tarjeta_a_efectivo_por_sobrante"
    assert r.evidencia["sobrante_efectivo"] == 30000 and r.evidencia["efectivo_contado"] == 50000
    assert "sobrante" in r.motivo
    assert not Excepcion.objects.filter(dia=dia, abierta=True).exists()


def test_la_correccion_deja_el_arqueo_cuadrado_sin_contar_dos_veces(caja_con_sobrante, usuario):
    dia, _ = caja_con_sobrante
    ejecutar_conciliacion(dia, usuario)
    arqueo = ArqueoEfectivo.objects.get(dia=dia, caja="CAJA 1")
    assert (arqueo.teorico, arqueo.ajustes, arqueo.esperado, arqueo.diferencia) == (20000, 30000, 50000, 0)


def test_la_regla_automatica_queda_en_la_bitacora_y_en_el_reporte(caja_con_sobrante, usuario):
    dia, falsa = caja_con_sobrante
    ejecutar_conciliacion(dia, usuario)
    evento = EventoAuditoria.objects.get(accion="reclasificacion_automatica")
    assert evento.valor_posterior["medio_corregido"] == "efectivo" and evento.valor_posterior["monto"] == "30000"

    (fila,) = filas_reclasificacion(dia)
    assert fila["nota_de_venta"] == 2 and fila["medio_original"] == "Débito" and fila["medio_corregido"] == "Efectivo"
    assert fila["automatica"] and fila["quien"].startswith("Regla automática")
    assert "Sobrante de efectivo: $30.000" in fila["evidencia"]


def test_sin_sobrante_la_venta_queda_como_excepcion_abierta(dia, usuario):
    _venta(dia, 1, "efectivo", 20000, time(10, 0))
    falsa = _venta(dia, 2, "debito", 30000, time(11, 0))
    _contar(dia, "CAJA 1", 20000)  # el efectivo cuadra: no sobra nada
    ejecutar_conciliacion(dia, usuario)
    assert not Reclasificacion.objects.exists()
    assert Excepcion.objects.filter(dia=dia, abierta=True, pago=falsa, tipo="pago_sin_movimiento").exists()


def test_sin_conteo_de_caja_no_se_corrige_nada(dia, usuario):
    _venta(dia, 2, "debito", 30000, time(11, 0))
    ejecutar_conciliacion(dia, usuario)
    assert not Reclasificacion.objects.exists()


def test_si_hay_dos_explicaciones_posibles_no_se_corrige_y_se_deja_una_pista(dia, usuario):
    _venta(dia, 1, "debito", 30000, time(10, 0))
    _venta(dia, 2, "debito", 30000, time(11, 0))
    _contar(dia, "CAJA 1", 30000)  # sobran $30.000 y podría ser cualquiera de las dos
    ejecutar_conciliacion(dia, usuario)
    assert not Reclasificacion.objects.exists()
    pistas = Excepcion.objects.filter(dia=dia, abierta=True, decision_propuesta__contains="Posible pago en efectivo")
    assert pistas.count() == 2


def test_la_regla_no_mira_otras_cajas(dia, usuario):
    _venta(dia, 1, "efectivo", 20000, time(10, 0), caja="CAJA 1")
    _venta(dia, 2, "debito", 30000, time(11, 0), caja="CAJA 2")
    _contar(dia, "CAJA 1", 50000)  # el sobrante es de la caja 1, la venta es de la caja 2
    ejecutar_conciliacion(dia, usuario)
    assert not Reclasificacion.objects.exists()


def test_reconciliar_de_nuevo_no_duplica_y_reevalua_con_el_conteo_vigente(caja_con_sobrante, usuario):
    dia, falsa = caja_con_sobrante
    ejecutar_conciliacion(dia, usuario)
    ejecutar_conciliacion(dia, usuario)
    assert Reclasificacion.objects.count() == 1

    _contar(dia, "CAJA 1", 20000)  # se recuenta: ya no sobra nada
    ejecutar_conciliacion(dia, usuario)
    falsa.refresh_from_db()
    assert falsa.medio_aprobado == "" and not Reclasificacion.objects.exists()
    assert Excepcion.objects.filter(dia=dia, abierta=True, pago=falsa).exists()


def test_al_registrar_el_conteo_de_un_dia_ya_conciliado_la_regla_se_aplica_sola(dia, usuario):
    _venta(dia, 1, "efectivo", 20000, time(10, 0))
    falsa = _venta(dia, 2, "debito", 30000, time(11, 0))
    ejecutar_conciliacion(dia, usuario)  # todavía sin conteo
    assert not Reclasificacion.objects.exists()

    arqueo = ArqueoEfectivo.objects.get(dia=dia, caja="CAJA 1")
    registrar_conteo(arqueo, fondo_inicial=Decimal(0), gastos_autorizados=Decimal(0),
                     total_contado=Decimal(50000), usuario=usuario)
    falsa.refresh_from_db()
    assert falsa.medio_aprobado == "efectivo"
    arqueo.refresh_from_db()
    assert arqueo.diferencia == 0


# ------------------------------------------------- efectivo -> tarjeta

@pytest.fixture
def caja_con_faltante(dia):
    """Efectivo $20.000 y $15.000; Transbank cobró $15.000 sin venta; la caja
    tiene solo $20.000: faltan $15.000."""
    _venta(dia, 1, "efectivo", 20000, time(10, 0))
    supuesta = _venta(dia, 2, "efectivo", 15000, time(10, 3))
    cobro = _cobro(dia, 15000, 10, 5, autorizacion="A9")
    _contar(dia, "CAJA 1", 20000)
    return dia, supuesta, cobro


def test_venta_declarada_como_efectivo_con_cobro_sin_venta_y_faltante_pasa_a_tarjeta(caja_con_faltante, usuario):
    dia, supuesta, cobro = caja_con_faltante
    ejecutar_conciliacion(dia, usuario)

    supuesta.refresh_from_db()
    assert (supuesta.medio, supuesta.medio_aprobado) == ("efectivo", "debito")
    r = Reclasificacion.objects.get(pago=supuesta)
    assert (r.medio_original, r.medio_corregido, r.monto, r.origen) == ("efectivo", "debito", 15000, "automatica")
    assert r.evidencia["autorizacion"] == "A9" and r.evidencia["movimiento_id"] == cobro.pk
    assert r.evidencia["faltante_efectivo"] == 15000
    assert Coincidencia.objects.filter(dia=dia, pagos=supuesta, movimientos=cobro).exists()
    assert not Excepcion.objects.filter(dia=dia, abierta=True).exists()
    arqueo = ArqueoEfectivo.objects.get(dia=dia, caja="CAJA 1")
    assert (arqueo.ajustes, arqueo.esperado, arqueo.diferencia) == (-15000, 20000, 0)


def test_con_faltante_pero_sin_venta_en_efectivo_del_mismo_monto_no_se_corrige(dia, usuario):
    _venta(dia, 1, "efectivo", 20000, time(10, 0))
    _cobro(dia, 15000, 10, 5)
    _contar(dia, "CAJA 1", 5000)  # faltan 15.000 pero no hay una venta en efectivo de 15.000
    ejecutar_conciliacion(dia, usuario)
    assert not Reclasificacion.objects.exists()
    assert Excepcion.objects.filter(dia=dia, abierta=True, tipo="movimiento_sin_pago").exists()


def test_con_dos_ventas_en_efectivo_igual_se_desempata_por_la_hora(dia, usuario):
    _venta(dia, 1, "efectivo", 15000, time(9, 0))
    cercana = _venta(dia, 2, "efectivo", 15000, time(15, 2))
    _cobro(dia, 15000, 15, 5)
    _contar(dia, "CAJA 1", 15000)  # faltan 15.000
    ejecutar_conciliacion(dia, usuario)
    cercana.refresh_from_db()
    assert cercana.medio_aprobado == "debito"
    assert Reclasificacion.objects.count() == 1


# ------------------------------------------------- correcciones manuales

def test_una_reclasificacion_manual_queda_con_usuario_motivo_y_evidencia(dia, usuario):
    pago = _venta(dia, 1, "debito", 50000, time(10, 0))
    cobro = _cobro(dia, 45000, 10, 1, autorizacion="A7")
    ejecutar_conciliacion(dia, usuario)
    excepcion = Excepcion.objects.get(dia=dia, pago=pago)
    motivo = MotivoResolucion.objects.get(codigo="mezcla_medios_pago")

    resolver_excepcion(
        excepcion, resolucion="mezcla_medios", motivo_id=motivo.pk, comentario="Pagó $5.000 en efectivo",
        usuario=usuario,
        lineas=[
            {"medio": "debito", "monto": 45000, "movimiento_id": cobro.pk},
            {"medio": "efectivo", "monto": 5000, "fuente_evidencia": "Confirmado en arqueo de caja"},
        ],
    )

    r = Reclasificacion.objects.get(pago=pago)  # solo la línea que cambia de medio
    assert (r.medio_original, r.medio_corregido, r.monto, r.origen, r.usuario) == ("debito", "efectivo", 5000, "manual", usuario)
    assert r.regla == "mezcla_de_medios" and r.motivo == motivo.texto
    assert r.evidencia["fuente"] == "Confirmado en arqueo de caja"
    # y el efectivo esperado de la caja sube por la parte pagada en efectivo
    assert calcular_ajustes_por_caja(dia) == {"CAJA 1": 5000}


def test_medio_contradictorio_aceptado_por_una_persona_deja_traza_con_la_autorizacion(dia, usuario):
    pago = _venta(dia, 1, "debito", 10000, time(10, 0))
    cobro = _cobro(dia, 10000, 10, 1, medio="credito", autorizacion="A5")
    ejecutar_conciliacion(dia, usuario)
    excepcion = Excepcion.objects.get(dia=dia, pago=pago, tipo="medio_contradictorio")
    motivo = MotivoResolucion.objects.filter(aplica_a="excepcion", activo=True).first()

    resolver_excepcion(excepcion, resolucion="aceptar", motivo_id=motivo.pk, comentario="ok", usuario=usuario)

    r = Reclasificacion.objects.get(pago=pago)
    assert (r.medio_original, r.medio_corregido, r.origen, r.usuario) == ("debito", "credito", "manual", usuario)
    assert r.evidencia["autorizacion"] == "A5" and r.evidencia["movimiento_id"] == cobro.pk


# ------------------------------------------------------- el reporte

def test_el_reporte_de_auditoria_se_ve_y_se_descarga(client, caja_con_sobrante, usuario, tmp_path, monkeypatch):
    dia, _ = caja_con_sobrante
    ejecutar_conciliacion(dia, usuario)
    client.force_login(usuario)

    resp = client.get(reverse("auditoria_reclasificaciones", args=["2026-09-01"]))
    html = resp.content.decode()
    assert resp.status_code == 200 and "Regla automática" in html and "30.000" in html

    import nucleo.auditoria_servicio as svc
    monkeypatch.setattr(svc.settings, "BASE_DIR", tmp_path)
    resp = client.get(reverse("descargar_auditoria", args=["2026-09-01"]))
    assert resp.status_code == 200
    ruta = generar_archivo_auditoria(dia, tmp_path)
    filas = list(load_workbook(ruta).active.iter_rows(values_only=True))
    assert filas[0][:2] == ("Nota de venta", "Comprobante") and len(filas) == 2
    assert filas[1][4:7] == ("Débito", "Efectivo", 30000)


def test_un_dia_sin_correcciones_lo_dice(client, dia, usuario):
    client.force_login(usuario)
    html = client.get(reverse("auditoria_reclasificaciones", args=["2026-09-01"])).content.decode()
    assert "no tiene correcciones" in html


def test_un_dia_que_no_existe_lleva_a_cargar_archivos(client, usuario, db):
    client.force_login(usuario)
    resp = client.get(reverse("auditoria_reclasificaciones", args=["2026-09-30"]))
    assert resp.status_code == 302
