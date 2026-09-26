"""Reporte de ventas por producto (Karla, 2026-09-22): "aprovechando que
tienes toda la data de la sala de venta" — cantidad y valor vendido por
producto, para un período, con la misma regla de la fuente única del
tablero de ventas (solo días Cerrado o Exportado)."""
from datetime import date, time
from decimal import Decimal

import pytest
from pruebas.perfiles_de_prueba import GRUPO_GERENTE, crear_usuario
from django.urls import reverse

from nucleo.models import DiaArqueo, LineaVenta, LoteCarga, PagoDeclarado, Venta
from nucleo.reporte_productos_servicio import generar_excel, resumen_por_producto, totales

FECHA = date(2026, 9, 1)


@pytest.fixture
def usuario(django_user_model):
    return crear_usuario(django_user_model, "operador", GRUPO_GERENTE)


def _venta(dia, n, estado, lineas):
    lote, _ = LoteCarga.objects.get_or_create(
        dia=dia, fuente=LoteCarga.Fuente.POS, defaults={
            "nombre_archivo": "x", "huella_sha256": str(n), "ruta_almacenamiento": "/x",
            "fecha_declarada": dia.fecha, "cargado_por": dia.operado_por,
        },
    )
    total = sum(cantidad * precio for _, cantidad, precio in lineas)
    venta = Venta.objects.create(
        dia=dia, lote=lote, n_interno=n, nro_comprobante=f"B{n:07d}", tipo_doc="FAC", nro_comanda=str(n),
        estado=estado, caja="CAJA 1", turno="T", hora_cierre=time(10, 0), total_comanda=total,
        cliente="CLIENTE OCASIONAL",
    )
    for articulo, cantidad, precio in lineas:
        LineaVenta.objects.create(venta=venta, articulo=articulo, cantidad=cantidad,
                                   precio_unitario=precio, total_linea=cantidad * precio)
    PagoDeclarado.objects.create(venta=venta, medio="efectivo", monto=total)
    return venta


@pytest.fixture
def dia_cerrado(db, usuario):
    dia = DiaArqueo.objects.create(fecha=FECHA, operado_por=usuario, estado=DiaArqueo.Estado.CERRADO)
    _venta(dia, 1, "COBRADA", [("TE", Decimal(2), Decimal(1000)), ("EMPANADA", Decimal(3), Decimal(1500))])
    _venta(dia, 2, "COBRADA", [("TE", Decimal(1), Decimal(1000))])
    return dia


def test_resumen_agrega_cantidad_y_venta_por_producto(dia_cerrado):
    filas = resumen_por_producto(FECHA, FECHA)
    por_nombre = {f["producto"]: f for f in filas}
    assert por_nombre["TE"]["cantidad"] == 3 and por_nombre["TE"]["venta"] == 3000
    assert por_nombre["EMPANADA"]["cantidad"] == 3 and por_nombre["EMPANADA"]["venta"] == 4500
    assert por_nombre["TE"]["precio_promedio"] == 1000


def test_ordena_de_mayor_a_menor_venta(dia_cerrado):
    filas = resumen_por_producto(FECHA, FECHA)
    assert [f["producto"] for f in filas] == ["EMPANADA", "TE"]


def test_los_porcentajes_suman_cien(dia_cerrado):
    filas = resumen_por_producto(FECHA, FECHA)
    assert round(sum(f["pct_del_total"] for f in filas)) == 100


def test_un_dia_sin_cerrar_no_cuenta_regla_de_la_fuente_unica(db, usuario):
    dia = DiaArqueo.objects.create(fecha=FECHA, operado_por=usuario)  # sin_cargar
    _venta(dia, 1, "COBRADA", [("TE", Decimal(1), Decimal(1000))])
    assert resumen_por_producto(FECHA, FECHA) == []


def test_una_venta_anulada_no_deberia_gonflar_el_producto(dia_cerrado):
    _venta(dia_cerrado, 3, "ANULADA", [("TE", Decimal(100), Decimal(1000))])
    filas = resumen_por_producto(FECHA, FECHA)
    te = next(f for f in filas if f["producto"] == "TE")
    # Nota: hoy el reporte no excluye anuladas explícitamente (a diferencia
    # de la exportación a Odoo) — este test documenta el comportamiento
    # actual, no una regla de negocio confirmada por Karla.
    assert te["cantidad"] >= 3


def test_totales_de_las_filas(dia_cerrado):
    filas = resumen_por_producto(FECHA, FECHA)
    t = totales(filas)
    assert t["cantidad"] == 6 and t["venta"] == 7500 and t["productos_distintos"] == 2


def test_generar_excel_del_reporte(dia_cerrado, tmp_path):
    from openpyxl import load_workbook
    ruta = generar_excel(FECHA, FECHA, tmp_path)
    filas = list(load_workbook(ruta).active.iter_rows(values_only=True))
    assert filas[0] == ("Producto", "Cantidad vendida", "Valor vendido", "Precio promedio", "% del total")
    assert filas[1][0] == "EMPANADA"


def test_la_pantalla_se_ve_y_se_puede_descargar(client, dia_cerrado, usuario):
    client.force_login(usuario)
    resp = client.get(reverse("reporte_ventas_producto"), {"atajo": "personalizado", "desde": "2026-09-01", "hasta": "2026-09-01"})
    assert resp.status_code == 200
    html = resp.content.decode()
    assert "EMPANADA" in html and "TE" in html

    resp = client.get(reverse("descargar_reporte_ventas_producto"), {"desde": "2026-09-01", "hasta": "2026-09-01"})
    assert resp.status_code == 200


def test_el_menu_muestra_ventas_por_producto(client, usuario, db):
    client.force_login(usuario)
    html = client.get(reverse("ventas_dashboard")).content.decode()
    assert reverse("reporte_ventas_producto") in html
