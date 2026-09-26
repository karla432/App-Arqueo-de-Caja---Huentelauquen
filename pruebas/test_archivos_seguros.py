"""Bloque A de seguridad (Karla, 2026-09-26): archivos que entran (Excel y comprobantes) y que salen (Excel).
Se prueba con archivos maliciosos de ejemplo, no solo con los buenos."""
import io
import zipfile
from datetime import date

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from nucleo import archivos_seguros as seg
from nucleo.archivos_seguros import ExcelInvalido, guardar_libro, neutralizar_libro, nombre_seguro, texto_seguro, validar_xlsx_subido
from nucleo.carga_servicio import ArchivoInvalido, _validar_archivo
from nucleo.comprobantes import ComprobanteInvalido, validar_comprobante


def _xlsx_bytes(celdas=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, v in enumerate(celdas or ["a", 1], start=1):
        ws.cell(row=i, column=1, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip(partes: dict[str, bytes], base: bytes | None = None) -> bytes:
    """Un .xlsx real al que se le agregan partes extra."""
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base or _xlsx_bytes())) as origen, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as destino:
        for info in origen.infolist():
            destino.writestr(info, origen.read(info.filename))
        for nombre, contenido in partes.items():
            destino.writestr(nombre, contenido)
    return buf.getvalue()


def _subido(nombre, contenido):
    return SimpleUploadedFile(nombre, contenido)


# --------------------------------------------------------------------------- Excel que entra

def test_un_xlsx_normal_pasa():
    validar_xlsx_subido(_subido("ventas.xlsx", _xlsx_bytes()))


def test_los_archivos_de_referencia_reales_pasan():
    from pathlib import Path
    carpeta = Path(__file__).parent / "fixtures"
    for ruta in carpeta.glob("*/*.xlsx"):
        validar_xlsx_subido(_subido(ruta.name, ruta.read_bytes()))


def test_una_extension_distinta_se_rechaza():
    with pytest.raises(ExcelInvalido, match="no permitido"):
        validar_xlsx_subido(_subido("ventas.xlsm", _xlsx_bytes()))
    with pytest.raises(ExcelInvalido, match="no permitido"):
        validar_xlsx_subido(_subido("ventas.csv", b"a,b"))


def test_un_archivo_renombrado_a_xlsx_se_rechaza():
    with pytest.raises(ExcelInvalido, match="no corresponde"):
        validar_xlsx_subido(_subido("ventas.xlsx", b"MZ\x90\x00 esto es un ejecutable"))
    with pytest.raises(ExcelInvalido, match="no corresponde"):
        validar_xlsx_subido(_subido("ventas.xlsx", b"%PDF-1.4 esto es un pdf"))


def test_un_zip_cualquiera_sin_estructura_de_excel_se_rechaza():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("cualquier.txt", "hola")
    with pytest.raises(ExcelInvalido, match="estructura"):
        validar_xlsx_subido(_subido("ventas.xlsx", buf.getvalue()))


def test_un_zip_dañado_se_rechaza():
    roto = _xlsx_bytes()[:-40]
    with pytest.raises(ExcelInvalido):
        validar_xlsx_subido(_subido("ventas.xlsx", roto))


def test_macros_y_objetos_incrustados_se_rechazan():
    for parte in ("xl/vbaProject.bin", "xl/embeddings/oleObject1.bin", "xl/activeX/activeX1.xml"):
        with pytest.raises(ExcelInvalido, match="macros"):
            validar_xlsx_subido(_subido("ventas.xlsx", _zip({parte: b"x"})))


def test_rutas_internas_peligrosas_se_rechazan():
    with pytest.raises(ExcelInvalido, match="rutas"):
        validar_xlsx_subido(_subido("ventas.xlsx", _zip({"../../evil.txt": b"x"})))


def test_una_bomba_de_compresion_se_rechaza(monkeypatch):
    monkeypatch.setattr(seg, "MAX_DESCOMPRIMIDO_MB", 1)
    bomba = _zip({"xl/relleno.bin": b"\0" * (3 * 1024 * 1024)})
    with pytest.raises(ExcelInvalido, match="se expande"):
        validar_xlsx_subido(_subido("ventas.xlsx", bomba))


def test_razon_de_compresion_anormal_se_rechaza():
    bomba = _zip({"xl/relleno.bin": b"\0" * (12 * 1024 * 1024)})
    with pytest.raises(ExcelInvalido, match="comprimido de forma anormal"):
        validar_xlsx_subido(_subido("ventas.xlsx", bomba))


def test_demasiadas_partes_se_rechaza():
    partes = {f"xl/extra/p{i}.xml": b"x" for i in range(seg.MAX_PARTES + 5)}
    with pytest.raises(ExcelInvalido, match="demasiadas"):
        validar_xlsx_subido(_subido("ventas.xlsx", _zip(partes)))


def test_la_pantalla_de_carga_rechaza_lo_malicioso_con_su_mensaje():
    with pytest.raises(ArchivoInvalido, match="macros"):
        _validar_archivo(_subido("ventas.xlsx", _zip({"xl/vbaProject.bin": b"x"})))


def test_una_formula_en_un_excel_cargado_se_lee_como_valor_no_se_ejecuta(tmp_path):
    """pandas abre con data_only=True: una fórmula maliciosa llega como su valor guardado (o vacío), jamás se evalúa."""
    import pandas as pd
    wb = openpyxl.Workbook()
    wb.active["A1"] = "dato"
    wb.active["A2"] = '=HYPERLINK("http://malo.example","clic")'
    ruta = tmp_path / "f.xlsx"
    wb.save(ruta)
    df = pd.read_excel(ruta, header=None)
    assert "HYPERLINK" not in "".join(str(x) for x in df[0].dropna()) or df[0].isna().any()


# --------------------------------------------------------------------------- Excel que sale

def test_los_textos_que_parecen_formula_se_neutralizan():
    assert texto_seguro('=HYPERLINK("http://malo.example","x")').startswith("'=")
    assert texto_seguro("@SUM(1+1)").startswith("'@")
    assert texto_seguro("+cmd|' /C calc'!A0").startswith("'+")
    assert texto_seguro("-2+3+cmd|' /C calc'!A0").startswith("'-")


def test_lo_normal_no_se_toca():
    for v in ("CLIENTE OCASIONAL", "Boleta 123", "-", "+", "-5000", "+12,5", "1.234", "-0,5", 15, 3.5, None, date(2026, 9, 1)):
        assert texto_seguro(v) == v


def test_un_libro_se_guarda_sin_formulas_inyectadas(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Cliente", "Monto"])
    ws.append(['=HYPERLINK("http://malo.example","x")', -1500])
    ws.append(["Comentario", "@SUM(A1)"])
    ruta = tmp_path / "salida.xlsx"
    guardar_libro(wb, ruta)
    leido = load_workbook(ruta)
    hoja = leido.active
    assert hoja["A2"].data_type != "f" and hoja["A2"].value.startswith("'=")
    assert hoja["B2"].value == -1500                       # los números no se tocan
    assert hoja["B3"].data_type != "f" and hoja["B3"].value.startswith("'@")
    assert hoja["A1"].value == "Cliente"


def test_neutralizar_cuenta_lo_que_toca():
    wb = openpyxl.Workbook()
    wb.active.append(["=1+1", "normal", "@x"])
    assert neutralizar_libro(wb) == 2


def test_el_reporte_exportado_por_la_app_no_lleva_formulas_de_datos_maliciosos(tmp_path, django_user_model):
    """De punta a punta: un artículo llamado como una fórmula no llega como fórmula al Excel del reporte."""
    from datetime import time
    from nucleo.models import DiaArqueo, LineaVenta, LoteCarga, PagoDeclarado, Venta
    from nucleo.reporte_productos_servicio import generar_excel
    u = django_user_model.objects.create_user(username="x", password="x")
    dia = DiaArqueo.objects.create(fecha=date(2026, 9, 2), operado_por=u, estado="cerrado")
    lote = LoteCarga.objects.create(dia=dia, fuente="pos", nombre_archivo="x", huella_sha256="a", ruta_almacenamiento="/x",
                                    fecha_declarada=dia.fecha, cargado_por=u)
    v = Venta.objects.create(dia=dia, lote=lote, n_interno=1, nro_comprobante="B1", tipo_doc="FAC", nro_comanda="1",
                             estado="COBRADA", caja="CAJA 1", turno="T", hora_cierre=time(10, 0), total_comanda=1000,
                             cliente="CLIENTE OCASIONAL")
    LineaVenta.objects.create(venta=v, articulo='=cmd|" /C calc"!A0', cantidad=1, precio_unitario=1000, total_linea=1000)
    PagoDeclarado.objects.create(venta=v, medio="efectivo", monto=1000)
    ruta = generar_excel(dia.fecha, dia.fecha, tmp_path)
    textos = []
    for ws in load_workbook(ruta).worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
                if isinstance(c.value, str) and "cmd" in c.value:
                    textos.append(c.value)
    assert textos and all(t.startswith("'=") for t in textos)


# --------------------------------------------------------------------------- nombres

def test_los_nombres_de_archivo_no_pueden_escapar_de_su_carpeta():
    for malo in ("../../etc/passwd", "..\\..\\windows\\system32\\x.xlsx", "/abs/ruta/x.pdf", "C:\\Users\\x\\y.pdf", "a/b/c.pdf"):
        limpio = nombre_seguro(malo)
        assert "/" not in limpio and "\\" not in limpio and ".." not in limpio.replace("...", "")
    assert nombre_seguro("Ventas 01-09 (final).xlsx") == "Ventas_01-09_final.xlsx"
    assert nombre_seguro("") == "archivo"


# --------------------------------------------------------------------------- comprobantes

def _png_real() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buf, "PNG")
    return buf.getvalue()


def _jpg_real() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buf, "JPEG")
    return buf.getvalue()


def test_imagenes_reales_pasan_y_falsas_no():
    validar_comprobante(_subido("x.png", _png_real()))
    validar_comprobante(_subido("x.jpg", _jpg_real()))
    with pytest.raises(ComprobanteInvalido, match="imagen"):
        validar_comprobante(_subido("x.png", b"\x89PNG\r\n\x1a\n" + b"basura"))
    with pytest.raises(ComprobanteInvalido, match="imagen"):
        validar_comprobante(_subido("x.jpg", b"\xff\xd8\xff\xe0" + b"basura"))


def test_una_imagen_gigante_en_pixeles_se_rechaza(monkeypatch):
    from nucleo import comprobantes
    monkeypatch.setattr(comprobantes, "MAX_PIXELES", 10)
    with pytest.raises(ComprobanteInvalido, match="imagen"):
        validar_comprobante(_subido("x.png", _png_real()))


@pytest.mark.parametrize("activo", [
    b"/JavaScript (app.alert(1))", b"/JS (x)", b"/Launch /Win", b"/EmbeddedFile", b"/RichMedia", b"/J#61vaScript (x)",
])
def test_un_pdf_con_elementos_activos_se_rechaza(activo):
    pdf = b"%PDF-1.4\n1 0 obj << /Type /Catalog /OpenAction << /S " + activo + b" >> >> endobj\n%%EOF"
    with pytest.raises(ComprobanteInvalido, match="activos"):
        validar_comprobante(_subido("x.pdf", pdf))


def test_un_pdf_normal_con_palabras_parecidas_pasa():
    pdf = b"%PDF-1.4\n1 0 obj << /Type /Catalog /Pages 2 0 R /Title (Comprobante JSON) >> endobj\n%%EOF"
    validar_comprobante(_subido("x.pdf", pdf))


# --------------------------------------------------------------------------- XML malicioso dentro del Excel

def test_un_xml_con_entidades_maliciosas_no_se_procesa(tmp_path):
    """«Billion laughs»: sin defusedxml, openpyxl expandiría las entidades hasta agotar la memoria."""
    import openpyxl.xml
    assert openpyxl.xml.DEFUSEDXML, "defusedxml debe estar instalado (requirements.txt)"
    bomba = (
        b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;">]>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>&lol2;</sheetData></worksheet>'
    )
    ruta = tmp_path / "bomba.xlsx"
    ruta.write_bytes(_zip({"xl/worksheets/sheet1.xml": bomba}))  # duplicada: la parte válida y la maliciosa
    with pytest.raises(Exception):
        load_workbook(ruta)
