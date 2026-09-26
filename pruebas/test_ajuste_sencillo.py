"""Ajuste de sencillo (Karla, 2026-09-21): diferencia de hasta $10, positiva o
negativa, se calza sola; identificada con el motivo "Ajuste de sencillo" y
contabilizada según el signo en una de las dos cuentas fijas del sistema
(pérdida o ganancia — "solo ocuparemos dos cuentas contables"). No sale en
el archivo de ventas: va agregada en el asiento de arqueo de caja (Fase 2).
"""
from datetime import date, datetime, time

import pytest
from pruebas.perfiles_de_prueba import GRUPO_GERENTE, crear_usuario
from openpyxl import load_workbook

from nucleo.cierre_servicio import aprobar_y_cerrar_dia
from nucleo.conciliacion_servicio import ejecutar_conciliacion
from nucleo.exportacion_servicio import generar_exportacion
from nucleo.arqueo_servicio import registrar_conteo
from nucleo.models import (
    ArqueoEfectivo, Coincidencia, DiaArqueo, Excepcion, LineaVenta, LoteCarga,
    MovimientoProveedor, PagoDeclarado, Venta,
)
from decimal import Decimal

FECHA = date(2026, 9, 1)


@pytest.fixture
def usuario(django_user_model):
    return crear_usuario(django_user_model, "operador", GRUPO_GERENTE)


@pytest.fixture
def dia_con_dos_diferencias(db, usuario):
    dia = DiaArqueo.objects.create(fecha=FECHA, operado_por=usuario)
    lote_pos = LoteCarga.objects.create(
        dia=dia, fuente=LoteCarga.Fuente.POS, nombre_archivo="x.xlsx", huella_sha256="a",
        ruta_almacenamiento="/tmp/x", fecha_declarada=FECHA, cargado_por=usuario,
    )
    lote_tbk = LoteCarga.objects.create(
        dia=dia, fuente=LoteCarga.Fuente.TRANSBANK, nombre_archivo="t.xlsx", huella_sha256="b",
        ruta_almacenamiento="/tmp/t", fecha_declarada=FECHA, cargado_por=usuario,
    )
    # POS $58.580 vs proveedor $58.579 (POS mayor) y POS $57.164 vs $57.165 (proveedor mayor)
    for n, (monto_pos, monto_tbk, hora) in enumerate([(58580, 58579, 12), (57164, 57165, 15)], start=1):
        venta = Venta.objects.create(
            dia=dia, lote=lote_pos, n_interno=n, nro_comprobante=f"B000000{n}", tipo_doc="FAC",
            nro_comanda=str(n), estado="COBRADA", caja="CAJA 1", turno="TURNO CAJA 1",
            hora_cierre=time(hora, 0), total_comanda=monto_pos, cliente="CLIENTE OCASIONAL",
        )
        LineaVenta.objects.create(venta=venta, articulo="X", cantidad=1, precio_unitario=monto_pos, total_linea=monto_pos)
        PagoDeclarado.objects.create(venta=venta, medio="debito", monto=monto_pos)
        MovimientoProveedor.objects.create(
            dia=dia, lote=lote_tbk, proveedor="transbank", tipo="Venta", estado="PENDIENTE",
            fecha_hora=datetime(2026, 9, 1, hora, 1), medio="debito", tipo_tarjeta="debito",
            monto_original=monto_tbk, monto_venta_valido_abono=monto_tbk, autorizacion=f"A{n}", terminal="T1",
        )
    return dia


def test_diferencias_de_sencillo_se_calzan_solas_sin_excepcion_abierta(dia_con_dos_diferencias, usuario):
    resultado = ejecutar_conciliacion(dia_con_dos_diferencias, usuario)
    assert len(resultado.coincidencias) == 2
    assert not Excepcion.objects.filter(dia=dia_con_dos_diferencias, abierta=True).exists()

    ajustes = Excepcion.objects.filter(
        dia=dia_con_dos_diferencias, resolucion=Excepcion.Resolucion.AJUSTE_SENCILLO,
    ).order_by("pk")
    assert ajustes.count() == 2
    for ajuste in ajustes:
        assert ajuste.motivo.codigo == "ajuste_sencillo"
        assert ajuste.motivo.texto == "Ajuste de sencillo"
        assert ajuste.abierta is False
        assert ajuste.monto_perdida_ganancia == 1  # siempre sin signo; el signo lo da la cuenta
    # Signo: POS − proveedor. +1 (el proveedor cobró menos) es pérdida;
    # −1 (cobró más) es ganancia — "solo ocuparemos dos cuentas, siempre
    # estas cuentas" (Karla, 2026-09-21).
    assert [a.cuenta_contable.codigo for a in ajustes] == ["6501004", "6501003"]


def test_reconciliar_de_nuevo_no_duplica_los_ajustes(dia_con_dos_diferencias, usuario):
    ejecutar_conciliacion(dia_con_dos_diferencias, usuario)
    ejecutar_conciliacion(dia_con_dos_diferencias, usuario)
    assert Excepcion.objects.filter(
        dia=dia_con_dos_diferencias, resolucion=Excepcion.Resolucion.AJUSTE_SENCILLO,
    ).count() == 2
    assert Coincidencia.objects.filter(dia=dia_con_dos_diferencias).count() == 2


def test_el_ajuste_de_sencillo_no_sale_en_el_archivo_de_ventas(
    dia_con_dos_diferencias, usuario, tmp_path, monkeypatch,
):
    """Karla, 2026-09-21: "Tipo de pago externo" y "Tipo de documento
    externo" son listas cerradas en Odoo — una cuenta contable no es un
    tipo de pago válido. El ajuste de sencillo se saca de este archivo y
    queda solo en el asiento de arqueo de caja (ver test_documentos_libro
    y arqueo_caja_servicio)."""
    import nucleo.exportacion_servicio as svc
    monkeypatch.setattr(svc, "CARPETA_EXPORTACIONES", tmp_path)
    ejecutar_conciliacion(dia_con_dos_diferencias, usuario)
    arqueo, _ = ArqueoEfectivo.objects.get_or_create(dia=dia_con_dos_diferencias, caja="CAJA 1")
    registrar_conteo(arqueo, fondo_inicial=Decimal(0), gastos_autorizados=Decimal(0),
                      total_contado=Decimal(0), usuario=usuario)
    aprobar_y_cerrar_dia(dia_con_dos_diferencias, usuario, motivo_excepcion="prueba")

    exportacion = generar_exportacion(dia_con_dos_diferencias, usuario)
    filas = list(load_workbook(exportacion.archivo).active.iter_rows(values_only=True))
    assert len(filas) == 3  # encabezado + las 2 boletas, nada de ajustes
    assert not any("AJUSTE" in str(c) for f in filas for c in f)


def test_dos_ajustes_que_se_cancelan_no_dejan_linea_en_el_asiento(dia_con_dos_diferencias, usuario):
    """+1 y −1 se anulan entre sí — no hay nada real que contabilizar."""
    from nucleo.arqueo_caja_servicio import generar_lineas_arqueo_caja

    ejecutar_conciliacion(dia_con_dos_diferencias, usuario)
    filas = generar_lineas_arqueo_caja(dia_con_dos_diferencias)
    assert not any(f["Apuntes contables/Cuenta"].startswith(("6501003", "6501004")) for f in filas)


@pytest.fixture
def dia_con_cuatro_ajustes(db, usuario):
    """Caso real del día 01/09/2026: tres diferencias de +$1 y una de −$1
    → neto +$2 (pérdida), tal como el ejemplo de asiento que entregó Karla."""
    dia = DiaArqueo.objects.create(fecha=FECHA, operado_por=usuario)
    lote_pos = LoteCarga.objects.create(
        dia=dia, fuente=LoteCarga.Fuente.POS, nombre_archivo="x.xlsx", huella_sha256="a",
        ruta_almacenamiento="/tmp/x", fecha_declarada=FECHA, cargado_por=usuario,
    )
    lote_tbk = LoteCarga.objects.create(
        dia=dia, fuente=LoteCarga.Fuente.TRANSBANK, nombre_archivo="t.xlsx", huella_sha256="b",
        ruta_almacenamiento="/tmp/t", fecha_declarada=FECHA, cargado_por=usuario,
    )
    diferencias = [(58580, 58579, 12), (25133, 25132, 13), (46819, 46818, 14), (57164, 57165, 15)]
    for n, (monto_pos, monto_tbk, hora) in enumerate(diferencias, start=1):
        venta = Venta.objects.create(
            dia=dia, lote=lote_pos, n_interno=n, nro_comprobante=f"B000000{n}", tipo_doc="FAC",
            nro_comanda=str(n), estado="COBRADA", caja="CAJA 1", turno="TURNO CAJA 1",
            hora_cierre=time(hora, 0), total_comanda=monto_pos, cliente="CLIENTE OCASIONAL",
        )
        LineaVenta.objects.create(venta=venta, articulo="X", cantidad=1, precio_unitario=monto_pos, total_linea=monto_pos)
        PagoDeclarado.objects.create(venta=venta, medio="debito", monto=monto_pos)
        MovimientoProveedor.objects.create(
            dia=dia, lote=lote_tbk, proveedor="transbank", tipo="Venta", estado="PENDIENTE",
            fecha_hora=datetime(2026, 9, 1, hora, 1), medio="debito", tipo_tarjeta="debito",
            monto_original=monto_tbk, monto_venta_valido_abono=monto_tbk, autorizacion=f"A{n}", terminal="T1",
        )
    return dia


def test_los_ajustes_de_sencillo_del_dia_salen_en_una_sola_linea_neta(dia_con_cuatro_ajustes, usuario):
    from nucleo.arqueo_caja_servicio import generar_lineas_arqueo_caja

    ejecutar_conciliacion(dia_con_cuatro_ajustes, usuario)
    filas = generar_lineas_arqueo_caja(dia_con_cuatro_ajustes)
    resultado = [f for f in filas if f["Apuntes contables/Cuenta"].startswith(("6501003", "6501004"))]
    assert len(resultado) == 1
    fila = resultado[0]
    assert fila["Apuntes contables/Cuenta"] == "6501004 Pérdida por diferencia de efectivo"
    assert fila["Apuntes contables/Empresa"] == "TRANSBANK"
    assert fila["Apuntes contables/Débito"] == 2
    assert fila["Apuntes contables/Etiqueta"] == f"Diferencia Transbank -Arqueo Caja Día {FECHA.strftime('%d/%m/%Y')}"
    assert fila["Apuntes contables/Distribución analítica"] == '{"31": 100.0}'
