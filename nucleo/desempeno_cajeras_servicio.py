"""Desempeño de cajeras (Karla, 2026-10-26: «un informe de las cajeras, para ver quién se equivoca más»).

Por cajera y período: días trabajados, comandas, faltantes y sobrantes de efectivo, días fuera de umbral,
errores de medio de pago, cobros de más o de menos con tarjeta, notas de crédito y reintegros pendientes.

La cajera de una comanda es la asignada al arqueo de su caja ese día (ArqueoEfectivo.cajera): la app no
recibe la cajera desde el POS. Una caja sin cajera asignada queda en la fila «Sin cajera asignada», para que
el hueco se vea y se corrija en Arqueo de Efectivo.

Qué cuenta como error (base del semáforo, cada 100 comandas):
* Día fuera de umbral: la caja cerró con diferencia sobre «umbral_diferencia_caja» (en más o en menos).
* Error de medio: boleta registrada con un medio y cobrada con otro (cada pago reclasificado, una vez).
* Cobro de más o de menos: excepción de tarjeta resuelta como pérdida o ganancia de caja.
Las notas de crédito se muestran pero no suman errores: una devolución no siempre es un error de la cajera.

Misma regla de la fuente única que el tablero de ventas: solo cuentan días Cerrado o Exportado.
"""
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from nucleo.archivos_seguros import guardar_libro
from nucleo.models import (
    ArqueoEfectivo, CuentaContable, DiaArqueo, Excepcion, Reclasificacion, RegularizacionDocumento, Venta,
)
from nucleo.reintegros_servicio import saldo_pendiente
from nucleo.umbrales import umbral
from nucleo.ventas_dashboard_servicio import ESTADOS_INCLUIDOS

SIN_CAJERA = "Sin cajera asignada"
TOP = 5


def limites_semaforo() -> tuple[Decimal, Decimal]:
    verde = umbral("semaforo_cajeras_verde")
    amarillo = max(umbral("semaforo_cajeras_amarillo"), verde)
    return verde, amarillo


def color_semaforo(tasa: float | None, verde: Decimal, amarillo: Decimal) -> str:
    if tasa is None:
        return "gris"
    if tasa <= float(verde):
        return "verde"
    if tasa <= float(amarillo):
        return "amarillo"
    return "rojo"


def _fila_vacia(nombre: str, cajera_id) -> dict:
    return {
        "cajera": nombre, "cajera_id": cajera_id, "dias": set(), "comandas": 0, "venta": Decimal(0),
        "faltante": Decimal(0), "sobrante": Decimal(0), "dias_fuera_umbral": 0, "errores_medio": 0,
        "cobros_de_mas": 0, "cobros_de_menos": 0, "monto_cobros": Decimal(0), "notas_credito": 0,
        "monto_nc": Decimal(0), "reintegro_pendiente": Decimal(0),
    }


def desempeno(desde: date, hasta: date) -> dict:
    tope = umbral("umbral_diferencia_caja")
    verde, amarillo = limites_semaforo()
    dias = DiaArqueo.objects.filter(fecha__range=(desde, hasta), estado__in=ESTADOS_INCLUIDOS)
    pendientes = DiaArqueo.objects.filter(fecha__range=(desde, hasta)).exclude(
        estado__in=(*ESTADOS_INCLUIDOS, DiaArqueo.Estado.SIN_CARGAR)).count()

    filas: dict = {}
    incidentes: list[dict] = []

    def fila(cajera) -> dict:
        clave = cajera.pk if cajera else None
        if clave not in filas:
            filas[clave] = _fila_vacia(cajera.nombre if cajera else SIN_CAJERA, clave)
        return filas[clave]

    def incidente(f, dia, caja, tipo, monto, detalle, comanda="", destino="arqueo_efectivo"):
        incidentes.append({
            "fecha": dia.fecha, "cajera": f["cajera"], "caja": caja, "tipo": tipo, "comanda": comanda,
            "monto": monto, "detalle": detalle, "destino": destino,
        })

    # Caja de un día -> cajera (la asignada en el arqueo).
    cajera_de_caja: dict = {}
    arqueos = ArqueoEfectivo.objects.filter(dia__in=dias).select_related("cajera", "dia")
    for a in arqueos:
        cajera_de_caja[(a.dia_id, a.caja)] = a.cajera
        f = fila(a.cajera)
        f["dias"].add(a.dia_id)
        if a.diferencia is None:
            continue
        if a.diferencia < 0:
            f["faltante"] += -a.diferencia
        elif a.diferencia > 0:
            f["sobrante"] += a.diferencia
        if abs(a.diferencia) > tope:
            f["dias_fuera_umbral"] += 1
            incidente(f, a.dia, a.caja, "Faltante de efectivo" if a.diferencia < 0 else "Sobrante de efectivo",
                      a.diferencia, f"Diferencia de caja sobre el umbral de ±${int(tope):,}".replace(",", "."))
        if a.cajera_id and a.diferencia < 0:
            f["reintegro_pendiente"] += max(saldo_pendiente(a), Decimal(0))

    def cajera_de(dia_id, caja):
        return cajera_de_caja.get((dia_id, caja))

    for v in Venta.objects.filter(dia__in=dias).only("dia_id", "caja", "total_comanda"):
        f = fila(cajera_de(v.dia_id, v.caja))
        f["comandas"] += 1
        f["venta"] += v.total_comanda
        f["dias"].add(v.dia_id)

    # Errores de medio: un pago reclasificado cuenta una vez, aunque tenga varias reclasificaciones.
    vistos = set()
    for r in (Reclasificacion.objects.filter(dia__in=dias).select_related("pago__venta", "dia")
              .order_by("pago_id", "creada_en")):
        if r.pago_id in vistos:
            continue
        vistos.add(r.pago_id)
        venta = r.pago.venta
        f = fila(cajera_de(venta.dia_id, venta.caja))
        f["errores_medio"] += 1
        incidente(f, r.dia, venta.caja, "Error de medio de pago", r.monto,
                  f"Registrada como {r.get_medio_original_display().lower()}, cobrada como "
                  f"{r.get_medio_corregido_display().lower()}", comanda=venta.n_interno, destino="auditoria_reclasificaciones")

    # Cobros de más o de menos con tarjeta.
    for e in (Excepcion.objects.filter(dia__in=dias, resolucion=Excepcion.Resolucion.PERDIDA_GANANCIA)
              .select_related("cajera_responsable", "cuenta_contable", "pago__venta", "dia")):
        venta = e.pago.venta if e.pago_id else None
        caja = venta.caja if venta else ""
        cajera = e.cajera_responsable or (cajera_de(venta.dia_id, venta.caja) if venta else None)
        f = fila(cajera)
        de_menos = bool(e.cuenta_contable_id and e.cuenta_contable.tipo == CuentaContable.Tipo.PERDIDA)
        f["cobros_de_menos" if de_menos else "cobros_de_mas"] += 1
        f["monto_cobros"] += e.monto_perdida_ganancia
        incidente(f, e.dia, caja, "Cobro de menos" if de_menos else "Cobro de más", e.monto_perdida_ganancia,
                  (e.comentario or e.get_tipo_display())[:140], comanda=venta.n_interno if venta else "",
                  destino="bandeja_excepciones")
        if de_menos and e.cajera_responsable_id:
            f["reintegro_pendiente"] += max(saldo_pendiente(e), Decimal(0))

    for n in RegularizacionDocumento.objects.filter(dia__in=dias, tipo=RegularizacionDocumento.Tipo.NOTA_CREDITO):
        f = fila(cajera_de(n.dia_id, n.caja))
        f["notas_credito"] += 1
        f["monto_nc"] += n.monto or 0

    resultado = []
    for f in filas.values():
        if not f["comandas"] and not f["dias"]:
            continue
        f["dias_trabajados"] = len(f["dias"])
        del f["dias"]
        f["errores"] = f["dias_fuera_umbral"] + f["errores_medio"] + f["cobros_de_mas"] + f["cobros_de_menos"]
        f["tasa"] = (f["errores"] / f["comandas"] * 100) if f["comandas"] else None
        f["semaforo"] = color_semaforo(f["tasa"], verde, amarillo) if f["cajera_id"] else "gris"
        resultado.append(f)
    resultado.sort(key=lambda f: (f["cajera_id"] is None, -(f["tasa"] or 0), -f["errores"], f["cajera"]))

    con_cajera = [f for f in resultado if f["cajera_id"]]
    top = [f for f in con_cajera if f["errores"]][:TOP]
    incidentes.sort(key=lambda i: (i["fecha"], i["cajera"], i["caja"]))
    totales = {
        clave: sum(f[clave] for f in resultado)
        for clave in ("comandas", "venta", "faltante", "sobrante", "dias_fuera_umbral", "errores_medio",
                      "cobros_de_mas", "cobros_de_menos", "notas_credito", "reintegro_pendiente", "errores")
    }
    totales["tasa"] = (totales["errores"] / totales["comandas"] * 100) if totales["comandas"] else None
    sin_cajera = next((f for f in resultado if f["cajera_id"] is None), None)
    return {
        "desde": desde, "hasta": hasta, "filas": resultado, "top": top, "incidentes": incidentes,
        "totales": totales, "verde": verde, "amarillo": amarillo, "umbral_caja": tope,
        "dias_incluidos": dias.count(), "dias_sin_cerrar": pendientes, "sin_cajera": sin_cajera,
    }


COLUMNAS = [
    "Cajera", "Días trabajados", "Comandas", "Venta", "Faltante", "Sobrante", "Días fuera de umbral",
    "Errores de medio", "Cobros de más", "Cobros de menos", "Total errores", "Errores c/100 comandas",
    "Semáforo", "Notas de crédito", "Reintegro pendiente",
]
RELLENO = {"verde": "E9F5EE", "amarillo": "FDF3E3", "rojo": "F9E0E4", "gris": "F2F2F2"}


def generar_excel(desde: date, hasta: date, carpeta: Path | None = None) -> Path:
    from nucleo.reporte_productos_servicio import CARPETA_REPORTES

    datos = desempeno(desde, hasta)
    carpeta = carpeta or CARPETA_REPORTES
    carpeta.mkdir(exist_ok=True)
    ruta = carpeta / f"desempeno_cajeras_{desde.isoformat()}_a_{hasta.isoformat()}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Desempeño de cajeras · {desde:%d-%m-%Y} al {hasta:%d-%m-%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Días cerrados incluidos: {datos['dias_incluidos']} · semáforo verde ≤ {int(datos['verde'])}, "
               f"amarillo ≤ {int(datos['amarillo'])}, rojo > {int(datos['amarillo'])} errores c/100 comandas"])
    ws.append([])
    ws.append(COLUMNAS)
    for celda in ws[4]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="ED0035")
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    for f in datos["filas"]:
        ws.append([
            f["cajera"], f["dias_trabajados"], f["comandas"], int(f["venta"]), int(f["faltante"]), int(f["sobrante"]),
            f["dias_fuera_umbral"], f["errores_medio"], f["cobros_de_mas"], f["cobros_de_menos"], f["errores"],
            round(f["tasa"], 2) if f["tasa"] is not None else None, f["semaforo"].capitalize(),
            f["notas_credito"], int(f["reintegro_pendiente"]),
        ])
        ws.cell(row=ws.max_row, column=13).fill = PatternFill("solid", fgColor=RELLENO[f["semaforo"]])
    for col in (4, 5, 6, 15):
        for (celda,) in ws.iter_rows(min_row=5, min_col=col, max_col=col):
            celda.number_format = '"$"#,##0'
    ws.column_dimensions["A"].width = 28
    for letra in "BCDEFGHIJKLMNO":
        ws.column_dimensions[letra].width = 14

    wi = wb.create_sheet("Incidentes")
    wi.append(["Fecha", "Cajera", "Caja", "Tipo", "Comanda", "Monto", "Detalle"])
    for celda in wi[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="ED0035")
    for i in datos["incidentes"]:
        wi.append([i["fecha"], i["cajera"], i["caja"], i["tipo"], i["comanda"] or None, int(i["monto"] or 0), i["detalle"]])
        wi.cell(row=wi.max_row, column=1).number_format = "DD-MM-YYYY"
        wi.cell(row=wi.max_row, column=6).number_format = '"$"#,##0'
    for letra, ancho in zip("ABCDEFG", (12, 26, 12, 24, 11, 12, 60)):
        wi.column_dimensions[letra].width = ancho
    guardar_libro(wb, ruta)
    return ruta
