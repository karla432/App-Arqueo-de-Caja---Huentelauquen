"""Consolidado por día y seguimiento de regularizaciones y reintegros (Karla,
2026-09-24; diseñado con los agentes UI, Contador Senior y Product Owner).

Dos preguntas que hasta hoy había que armar a mano recorriendo pantallas:

* **Consolidado por día** — una fila por día con lo vendido (boletas), cómo se
  cobró, el efectivo contado, las notas de crédito y guías, las excepciones, el
  depósito y las pérdidas de cajeras, con un cuadre «Sin explicar» que debe dar 0.
* **Seguimiento** — qué guías/notas siguen sin regularizar y cuánto debe cada
  cajera de sus pérdidas, con antigüedad para saber qué atender primero.

No calcula nada nuevo: agrega lo que ya existe. Las notas de crédito se muestran
solo como informativas (Resto ya las trae descontadas de las ventas y del efectivo).
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.conf import settings

from nucleo.dashboard_servicio import cuadre_por_medio
from nucleo.models import (
    ArqueoEfectivo, DepositoBancario, DiaArqueo, Excepcion, PagoDeclarado, Reclasificacion, RegularizacionDocumento, ReintegroCaja, Venta,
)
from nucleo.reintegros_servicio import (
    _es_perdida_de_arqueo, _es_perdida_de_excepcion, _reintegros_de, monto_perdida_original, saldo_pendiente,
)
from nucleo.umbrales import umbral

CARPETA_EXPORTACIONES = Path(settings.BASE_DIR) / "exportaciones"
CERO = Decimal(0)


def _suma(valores) -> Decimal:
    return sum((v for v in valores if v is not None), CERO)


def rango_del_mes(anio: int, mes: int) -> tuple[date, date]:
    siguiente = date(anio + (mes == 12), mes % 12 + 1, 1)
    return date(anio, mes, 1), date.fromordinal(siguiente.toordinal() - 1)


# ---------------------------------------------------------------- consolidado

def _perdidas_de_cajeras_del_dia(dia: DiaArqueo) -> tuple[Decimal, Decimal]:
    """(pérdidas atribuidas a una cajera, reintegrado validado) de ese día."""
    origenes = [a for a in dia.arqueos_caja.all() if _es_perdida_de_arqueo(a)]
    origenes += [
        e for e in Excepcion.objects.filter(dia=dia, resolucion=Excepcion.Resolucion.PERDIDA_GANANCIA)
        .select_related("cuenta_contable", "cajera_responsable") if _es_perdida_de_excepcion(e)
    ]
    perdido = _suma(monto_perdida_original(o) for o in origenes)
    reintegrado = _suma(monto_perdida_original(o) - saldo_pendiente(o) for o in origenes)
    return perdido, reintegrado


def fila_consolidada(dia: DiaArqueo) -> dict:
    ventas = Venta.objects.filter(dia=dia).exclude(estado="ANULADA")
    pagos = PagoDeclarado.objects.filter(venta__in=ventas)
    efectivo_pos = _suma(p.monto for p in pagos if p.medio == PagoDeclarado.Medio.EFECTIVO)
    tarjetas_pos = _suma(p.monto for p in pagos if p.medio in (
        PagoDeclarado.Medio.DEBITO, PagoDeclarado.Medio.CREDITO,
    ))
    boletas_total = _suma(v.total_comanda for v in ventas)
    otros_pos = _suma(p.monto for p in pagos) - efectivo_pos - tarjetas_pos

    docs = RegularizacionDocumento.objects.filter(dia=dia)
    guias = [d for d in docs if d.tipo == RegularizacionDocumento.Tipo.GUIA]
    notas = [d for d in docs if d.tipo == RegularizacionDocumento.Tipo.NOTA_CREDITO]
    sin_clasificar = [d for d in docs if d.tipo == RegularizacionDocumento.Tipo.SIN_CLASIFICAR]

    arqueos = list(dia.arqueos_caja.all())
    contados = [a for a in arqueos if a.contado is not None]
    excepciones = Excepcion.objects.filter(dia=dia)
    abiertas = [e for e in excepciones if e.abierta and e.tipo != "calidad_registro"]
    resultado_tarjeta = CERO
    for e in excepciones.filter(resolucion=Excepcion.Resolucion.PERDIDA_GANANCIA).select_related("cuenta_contable"):
        if e.cuenta_contable_id:
            signo = -1 if e.cuenta_contable.tipo == "perdida" else 1
            resultado_tarjeta += signo * e.monto_perdida_ganancia

    correcciones = list(Reclasificacion.objects.filter(dia=dia))  # correcciones de forma de pago (antes: menú «Auditoría de pagos»)
    depositos = [d for d in DepositoBancario.objects.filter(dia=dia) if d.monto_depositado is not None]
    perdido_cajeras, reintegrado = _perdidas_de_cajeras_del_dia(dia)
    sin_explicar = boletas_total - efectivo_pos - tarjetas_pos - otros_pos
    # Tarjetas contra Transbank (Karla, 2026-09-25, día 07): el cobro con tarjeta de una guía de despacho
    # está en el proveedor pero no en las ventas del POS; sin sumarlo, «tarjetas» nunca calza con Transbank.
    cuadre = {f["medio"]: f for f in cuadre_por_medio(dia)}
    guias_tarjeta = Decimal(str(sum(cuadre[m]["guias"] for m in ("debito", "credito"))))
    proveedor_tarjeta = Decimal(str(sum(cuadre[m]["proveedor"] or 0 for m in ("debito", "credito"))))
    tarjetas_total = tarjetas_pos + guias_tarjeta

    return {
        "dia": dia, "fecha": dia.fecha, "estado": dia.get_estado_display(),
        "boletas_n": ventas.count(), "boletas_total": boletas_total,
        "efectivo_pos": efectivo_pos, "tarjetas_pos": tarjetas_pos, "otros_pos": otros_pos, "sin_explicar": sin_explicar,
        "guias_tarjeta": guias_tarjeta, "tarjetas_total": tarjetas_total, "proveedor_tarjeta": proveedor_tarjeta,
        "diferencia_tarjeta": tarjetas_total - proveedor_tarjeta,
        "efectivo_esperado": _suma(a.esperado for a in contados) if contados else None,
        "efectivo_contado": _suma(a.contado for a in contados) if contados else None,
        "cajas_sin_contar": len(arqueos) - len(contados),
        "diferencia_efectivo": _suma(a.diferencia for a in contados) if contados else None,
        "guias_n": len(guias), "guias_total": _suma(d.monto for d in guias),
        "guias_efectivo": _suma(d.monto for d in guias if d.forma_pago == PagoDeclarado.Medio.EFECTIVO),
        "guias_pendientes": sum(1 for d in guias if d.estado == RegularizacionDocumento.Estado.PENDIENTE),
        "notas_n": len(notas), "notas_total": _suma(d.monto for d in notas),
        "notas_sin_respaldo": sum(1 for d in notas if d.verificacion_estado != RegularizacionDocumento.Verificacion.VERIFICADA),
        "sin_clasificar_n": len(sin_clasificar),
        "excepciones_abiertas_n": len(abiertas), "excepciones_abiertas_monto": _suma(e.monto_en_juego for e in abiertas),
        "resultado_tarjeta": resultado_tarjeta,
        "correcciones_pago_n": len(correcciones), "correcciones_pago_monto": _suma(r.monto for r in correcciones),
        "deposito_esperado": _suma(d.monto_esperado for d in depositos) if depositos else None,
        "deposito_depositado": _suma(d.monto_depositado for d in depositos) if depositos else None,
        "deposito_diferencia": _suma(d.diferencia for d in depositos) if depositos else None,
        "deposito_pendiente": DepositoBancario.objects.filter(dia=dia).exclude(estado__in=(
            DepositoBancario.Estado.APROBADO, DepositoBancario.Estado.OBSERVADO)).exists(),
        "perdidas_cajeras": perdido_cajeras, "reintegrado": reintegrado, "saldo_cajeras": perdido_cajeras - reintegrado,
        "cuadra": sin_explicar == 0 and not abiertas,
    }


def consolidado_por_dia(desde: date, hasta: date) -> list[dict]:
    dias = DiaArqueo.objects.filter(fecha__gte=desde, fecha__lte=hasta).exclude(
        estado=DiaArqueo.Estado.SIN_CARGAR).prefetch_related("arqueos_caja").order_by("fecha")
    return [fila_consolidada(d) for d in dias]


CAMPOS_SUMABLES = [
    "boletas_n", "boletas_total", "efectivo_pos", "tarjetas_pos", "otros_pos", "sin_explicar",
    "guias_tarjeta", "tarjetas_total", "proveedor_tarjeta", "diferencia_tarjeta", "efectivo_esperado",
    "efectivo_contado", "diferencia_efectivo", "guias_n", "guias_total", "guias_efectivo", "guias_pendientes",
    "notas_n", "notas_total", "notas_sin_respaldo", "sin_clasificar_n", "excepciones_abiertas_n",
    "excepciones_abiertas_monto", "resultado_tarjeta", "correcciones_pago_n", "correcciones_pago_monto", "deposito_esperado", "deposito_depositado",
    "deposito_diferencia", "perdidas_cajeras", "reintegrado", "saldo_cajeras",
]


def totales_consolidado(filas: list[dict]) -> dict:
    return {c: _suma(f[c] for f in filas) for c in CAMPOS_SUMABLES} | {
        "dias": len(filas), "dias_que_no_cuadran": sum(1 for f in filas if not f["cuadra"]),
    }


COLUMNAS_CONSOLIDADO = [
    ("Fecha", "fecha"), ("Estado del día", "estado"),
    ("Boletas (N)", "boletas_n"), ("Boletas ($)", "boletas_total"),
    ("Cobrado en efectivo (POS)", "efectivo_pos"), ("Cobrado con tarjeta (POS)", "tarjetas_pos"), ("Otros medios (POS)", "otros_pos"),
    ("Sin explicar (boletas − cobros)", "sin_explicar"),
    ("Guías cobradas con tarjeta", "guias_tarjeta"), ("Total tarjetas (POS + guías)", "tarjetas_total"),
    ("Transbank / proveedor", "proveedor_tarjeta"), ("Diferencia tarjetas", "diferencia_tarjeta"),
    ("Efectivo esperado", "efectivo_esperado"), ("Efectivo contado", "efectivo_contado"),
    ("Diferencia de efectivo", "diferencia_efectivo"),
    ("Guías (N)", "guias_n"), ("Guías ($)", "guias_total"), ("Guías cobradas en efectivo ($)", "guias_efectivo"),
    ("Guías por regularizar (N)", "guias_pendientes"),
    ("Notas de crédito (N) · informativo", "notas_n"), ("Notas de crédito ($) · informativo", "notas_total"),
    ("Notas sin respaldo (N)", "notas_sin_respaldo"),
    ("Excepciones abiertas (N)", "excepciones_abiertas_n"), ("Excepciones abiertas ($)", "excepciones_abiertas_monto"),
    ("Resultado por diferencias de tarjeta", "resultado_tarjeta"),
    ("Correcciones de forma de pago (N)", "correcciones_pago_n"), ("Correcciones de forma de pago ($)", "correcciones_pago_monto"),
    ("Depósito esperado", "deposito_esperado"), ("Depósito depositado", "deposito_depositado"),
    ("Diferencia de depósito", "deposito_diferencia"),
    ("Pérdidas atribuidas a cajeras", "perdidas_cajeras"), ("Reintegrado (validado)", "reintegrado"),
    ("Saldo por cobrar a cajeras", "saldo_cajeras"),
]


def generar_excel_consolidado(desde: date, hasta: date, carpeta: Path | None = None) -> Path:
    filas = consolidado_por_dia(desde, hasta)
    carpeta = carpeta or CARPETA_EXPORTACIONES
    carpeta.mkdir(exist_ok=True)
    ruta = carpeta / f"consolidado_por_dia_{desde.isoformat()}_{hasta.isoformat()}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado por día"
    ws.append([c for c, _ in COLUMNAS_CONSOLIDADO])
    for f in filas:
        ws.append([f[k] for _, k in COLUMNAS_CONSOLIDADO])
    if filas:
        t = totales_consolidado(filas)
        ws.append([])
        ws.append(["TOTAL", f"{t['dias']} días"] + [t.get(k) for _, k in COLUMNAS_CONSOLIDADO[2:]])
    else:
        ws.append([f"Sin días cargados entre {desde:%d/%m/%Y} y {hasta:%d/%m/%Y}"])
    wb.save(ruta)
    return ruta


# ----------------------------------------------------------------- seguimiento

def _antiguedad(fecha: date, hoy: date | None = None) -> int:
    return ((hoy or date.today()) - fecha).days


def _rango_antiguedad(dias: int) -> str:
    return "reciente" if dias <= 7 else "atencion" if dias <= 15 else "critico"


def seguimiento_documentos(desde: date, hasta: date, *, estado: str = "", tipo: str = "", hoy: date | None = None) -> list[dict]:
    """Guías, notas de crédito y códigos sin clasificar, con su antigüedad si siguen pendientes."""
    docs = RegularizacionDocumento.objects.filter(dia__fecha__gte=desde, dia__fecha__lte=hasta).select_related("dia").order_by("dia__fecha", "n_interno")
    if estado:
        docs = docs.filter(estado=estado)
    if tipo:
        docs = docs.filter(tipo=tipo)
    filas = []
    for d in docs:
        pendiente = d.estado == RegularizacionDocumento.Estado.PENDIENTE
        dias = _antiguedad(d.dia.fecha, hoy) if pendiente else None
        filas.append({
            "doc": d, "fecha": d.dia.fecha, "tipo": d.get_tipo_display(), "codigo": d.tipo, "folio": d.folio, "caja": d.caja,
            "monto": d.monto, "forma_pago": d.get_forma_pago_display() if d.forma_pago else "",
            "estado": d.get_estado_display(), "pendiente": pendiente,
            "relacionado": d.documento_relacionado, "dias_pendiente": dias,
            "nivel": _rango_antiguedad(dias) if dias is not None else "",
            "respaldo": d.get_verificacion_estado_display() if d.tipo == RegularizacionDocumento.Tipo.NOTA_CREDITO else "",
        })
    return filas


def _origenes_de_perdida(desde: date, hasta: date) -> list:
    arqueos = ArqueoEfectivo.objects.filter(dia__fecha__gte=desde, dia__fecha__lte=hasta).exclude(diferencia=None).select_related("cajera", "dia")
    excepciones = Excepcion.objects.filter(
        dia__fecha__gte=desde, dia__fecha__lte=hasta, resolucion=Excepcion.Resolucion.PERDIDA_GANANCIA,
    ).select_related("cajera_responsable", "cuenta_contable", "dia")
    return [a for a in arqueos if _es_perdida_de_arqueo(a)] + [e for e in excepciones if _es_perdida_de_excepcion(e)]


def seguimiento_reintegros(desde: date, hasta: date, *, estado: str = "", hoy: date | None = None) -> list[dict]:
    """Cada pérdida atribuida a una cajera con lo reintegrado (validado) y su saldo. Incluye las
    menores al umbral del panel (con marca) para que no se pierda trazabilidad."""
    minimo = umbral("umbral_reintegro_minimo")
    filas = []
    for o in _origenes_de_perdida(desde, hasta):
        original = monto_perdida_original(o)
        saldo = saldo_pendiente(o)
        estado_fila = "reintegrado" if saldo <= 0 else ("parcial" if saldo < original else "pendiente")
        if estado and estado != estado_fila:
            continue
        cajera = o.cajera if isinstance(o, ArqueoEfectivo) else o.cajera_responsable
        abonos = _reintegros_de(o)
        dias = _antiguedad(o.dia.fecha, hoy) if saldo > 0 else None
        filas.append({
            "origen": o, "tipo": "arqueo" if isinstance(o, ArqueoEfectivo) else "excepcion", "origen_id": o.pk,
            "fecha": o.dia.fecha, "cajera": cajera.nombre, "caja": o.caja if isinstance(o, ArqueoEfectivo) else "tarjeta",
            "original": original, "reintegrado": original - saldo, "saldo": saldo, "estado": estado_fila,
            "abonos_en_tramite": abonos.exclude(estado=ReintegroCaja.Estado.VALIDADO).count(),
            "menor": original <= minimo, "dias_pendiente": dias,
            "nivel": _rango_antiguedad(dias) if dias is not None else "",
        })
    return sorted(filas, key=lambda f: (f["estado"] == "reintegrado", f["fecha"]))


def resumen_por_cajera(filas: list[dict]) -> list[dict]:
    por_cajera: dict[str, dict] = {}
    for f in filas:
        r = por_cajera.setdefault(f["cajera"], {"cajera": f["cajera"], "original": CERO, "reintegrado": CERO, "saldo": CERO, "pendientes": 0})
        r["original"] += f["original"]
        r["reintegrado"] += f["reintegrado"]
        r["saldo"] += f["saldo"]
        r["pendientes"] += 1 if f["saldo"] > 0 else 0
    return sorted(por_cajera.values(), key=lambda r: -r["saldo"])


def generar_excel_seguimiento(desde: date, hasta: date, carpeta: Path | None = None) -> Path:
    documentos = seguimiento_documentos(desde, hasta)
    reintegros = seguimiento_reintegros(desde, hasta)
    carpeta = carpeta or CARPETA_EXPORTACIONES
    carpeta.mkdir(exist_ok=True)
    ruta = carpeta / f"seguimiento_regularizaciones_{desde.isoformat()}_{hasta.isoformat()}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append(["Fecha del día", "Tipo", "Folio", "Caja", "Monto", "Forma de pago", "Estado", "Documento relacionado", "Días pendiente", "Respaldo (notas)"])
    for f in documentos:
        ws.append([f["fecha"], f["tipo"], f["folio"], f["caja"], f["monto"], f["forma_pago"], f["estado"], f["relacionado"], f["dias_pendiente"], f["respaldo"]])
    if not documentos:
        ws.append(["Sin documentos en el período"])
    ws2 = wb.create_sheet("Pérdidas de cajeras")
    ws2.append(["Fecha del arqueo", "Cajera", "Origen", "Pérdida original", "Reintegrado (validado)", "Saldo pendiente", "Estado", "Días pendiente", "Menor al umbral"])
    for f in reintegros:
        ws2.append([f["fecha"], f["cajera"], f["caja"], f["original"], f["reintegrado"], f["saldo"], f["estado"], f["dias_pendiente"], "Sí" if f["menor"] else "No"])
    if not reintegros:
        ws2.append(["Sin pérdidas de cajeras en el período"])
    else:
        ws2.append([])
        ws2.append(["RESUMEN POR CAJERA"])
        for r in resumen_por_cajera(reintegros):
            ws2.append([None, r["cajera"], f"{r['pendientes']} pendientes", r["original"], r["reintegrado"], r["saldo"]])
    wb.save(ruta)
    return ruta
